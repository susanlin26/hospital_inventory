"""需要改善即轉改善追蹤，前線更新／提交，會計覆核結案及退回。"""
from io import BytesIO
from openpyxl import load_workbook
from test_workflow import client, sign, act, item, listing, PLAN
from app import monthly_cases as mc


def test_immediate_follow_and_frontline_can_progress_before_review(client):
 sign(client,'front099');act(client,'reason',**PLAN)
 t,k,x=item(client);assert x['display_status']=='待審核' and not x.get('cause_approved')
 assert mc.in_view(t,x,'follow','2026-09') and not mc.in_view(t,x,'confirm','2026-09')
 act(client,'progress',evidence='進度單001')
 x=item(client)[2];assert x['progress_note']=='測試查核說明' and not x.get('cause_approved')
 for view,n in [('confirm',1),('follow',2)]:
  r=client.get('/api/monthly-cases-export',params=dict(hospitals='SM',month='2026-09',baseline='mean3',domain='usage',view=view));assert r.status_code==200
  assert load_workbook(BytesIO(r.content))['案件明細'].max_row==n


def test_submit_before_reason_review_then_accountant_reject_finish(client):
 sign(client,'front099');act(client,'finish',expected=403)
 act(client,'submit',result_period='2026-09',evidence='已調整後領用單')
 t,k,x=item(client);assert x['display_status']=='待確認結案' and not x.get('cause_approved') and mc.in_view(t,x,'follow','2026-09')
 sign(client,'acct099');act(client,'progress',expected=403);act(client,'reject')
 t,k,x=item(client);assert x['return_stage']=='result' and mc.in_view(t,x,'follow','2026-09')
 sign(client,'front099');act(client,'submit',result_period='2026-09',evidence='完整領用核對單')
 sign(client,'acct099');act(client,'finish')
 t,k,x=item(client);assert x['display_status']=='已完成' and x['cause_approved'] and x['follow_up_completed']
 assert not mc.in_view(t,x,'follow','2026-09') and mc.in_view(t,x,'completed','2026-09')
 assert x['approved_cause']['category']==PLAN['cause_category'] and x['verified_data']


def test_plan_review_and_reason_return_stay_in_follow(client):
 sign(client,'front099');act(client,'reason',month='2026-10',**PLAN)
 sign(client,'acct099');act(client,'reject',month='2026-10')
 t,k,x=item(client,month='2026-10');assert mc.in_view(t,x,'follow','2026-10')
 sign(client,'front099');act(client,'progress',month='2026-10',expected=422);act(client,'submit',month='2026-10',expected=422,result_period='2026-10',evidence='待補正')
 act(client,'reason',month='2026-10',**PLAN)
 sign(client,'acct099');act(client,'approve_follow',month='2026-10')
 sign(client,'front099');act(client,'progress',month='2026-10',evidence='改善紀錄')
 assert item(client,month='2026-10')[2]['display_status']=='改善中'


def test_no_improvement_remains_confirm_and_viewer_cannot_write(client):
 sign(client,'front099');act(client,'reason',kind='idle',code='I',follow_up='no',cause_category='安全備貨')
 t,k,x=item(client,kind='idle',code='I');assert mc.in_view(t,x,'confirm','2026-09') and not mc.in_view(t,x,'follow','2026-09')
 sign(client,'view099');act(client,'progress',month='2026-10',expected=403)


def test_batch_reason_needing_improvement_enters_follow(client):
 sign(client,'front099')
 targets=[item(client,kind='purchase',code='U'),item(client,kind='high',code='L')]
 operations=[dict(case_id=t['id'],revision=t['revision'],keys=[k],action='reason',note='需要調整並持續追蹤',cause_category='其他',follow_up='yes',deadline='2026-11-30',plan='逐週核對並調整') for t,k,x in targets]
 r=client.post('/api/monthly-cases-batch',json=dict(operations=operations));assert r.status_code==200,r.text
 for kind,code in [('purchase','U'),('high','L')]:
  t,k,x=item(client,kind=kind,code=code);assert mc.in_view(t,x,'follow','2026-09') and not mc.in_view(t,x,'confirm','2026-09')
