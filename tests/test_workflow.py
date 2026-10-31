"""六種進度的端到端權限、證據保留、連結及統計驗證。使用隔離資料庫。"""
import os, json, uuid, subprocess
from pathlib import Path
from copy import deepcopy
from datetime import datetime
from decimal import Decimal as D
from io import BytesIO
import pytest
from fastapi.testclient import TestClient
from fastapi import HTTPException
from openpyxl import load_workbook
from psycopg.types.json import Jsonb
from app import db, auth, excel, monthly_workflow as wf, monthly_cases as mc

PW='Workflow-test-2026!'
ROOT=Path(__file__).parents[1]

@pytest.fixture(scope='module')
def client():
    assert os.environ['DATABASE_URL'].split('/')[-1].endswith('_test')
    import psycopg
    with psycopg.connect(db.dsn()) as c:
        c.execute('DROP SCHEMA public CASCADE'); c.execute('CREATE SCHEMA public')
    from app.server import app
    with TestClient(app,headers={'Origin':'http://testserver'}) as c:
        assert c.post('/api/setup',json=dict(username='admin099',name='測試管理員',password=PW)).status_code==200
        sign(c,'admin099')
        for name,role,h in [('front099','frontline','SM'),('acct099','accountant','SM'),('view099','viewer','SM'),('other099','frontline','DE')]:
            assert c.post('/api/users',json=dict(username=name,name=name,role=role,hospitals=[h],password=PW)).status_code==200
        with db.connect() as conn:
            uid=conn.execute("SELECT id FROM users WHERE username='admin099'").fetchone()['id']
            for month in range(4,11):
                p=f'2026-{month:02}'; fid=uuid.uuid4()
                conn.execute('INSERT INTO files(id,filename,sha256,content) VALUES(%s,%s,%s,%s)',(fid,f'{p}_SM測試.xlsx',str(fid),b'synthetic'))
                bid=conn.execute("INSERT INTO batches(hospital,period,version,status,active,file_id,file_hash,parser_version,report,operator,department) VALUES('SM',%s,1,'valid',true,%s,%s,'test',%s,%s,'RCW') RETURNING id",(p,fid,str(fid),Jsonb(dict(valid_count=3,classification_policy=excel.CLASSIFICATION_POLICY,tolerance_policy=excel.TOLERANCE_POLICY)),uid)).fetchone()['id']
                from helpers import line
                for i,(code,usage,qty,end) in enumerate([('U',10000 if month<9 else 30000 if month==9 else 40000,100,20000),('L',5000,1000,10),('I',0,0,5000)],2):
                    row=line(code=code,name='測試'+code,category='呼吸照護',period=p,sale_a=str(usage),sale_q=str(qty),open_a='50000',end_a=str(end),open_q='1000',end_q='10' if code=='L' else '200',in_a='1000' if month<9 else '5000',in_q='10',raw={'床別':'RCW'},row=i)
                    conn.execute('INSERT INTO lines VALUES(%s,%s,%s)',(bid,i,Jsonb(row)))
                conn.execute("INSERT INTO bed_departments(hospital,period,department,value,source,updated_by,measurement) VALUES('SM',%s,'RCW',%s,'test',%s,'patient_days')",(p,1000 if month<9 else 1200,uid))
        yield c

def sign(c,name):
    r=c.post('/api/login',json=dict(username=name,password=PW)); assert r.status_code==200,r.text
    c.headers['X-CSRF-Token']=r.json()['csrf']

def listing(c,month='2026-09',baseline='mean3'):
    r=c.get('/api/monthly-cases',params=dict(hospitals='SM',month=month,baseline=baseline)); assert r.status_code==200,r.text
    return r.json()

def item(c,kind='usage',month='2026-09',code='U',baseline='mean3'):
    data=listing(c,month,baseline)
    t=next(t for t in data['cases'] if t['kind']==kind and t['period']==month)
    k=next(k for k,x in t['items'].items() if x['data']['code']==code)
    return t,k,t['items'][k]

def act(c,action,kind='usage',month='2026-09',code='U',expected=200,**values):
    t,k,x=item(c,kind,month,code)
    r=c.post('/api/monthly-cases/'+str(t['id']),json=dict(revision=t['revision'],keys=[k],action=action,note='測試查核說明',**values));assert r.status_code==expected,r.text
    return r

PLAN=dict(follow_up='yes',owner='護理主管',deadline='2026-10-20',plan='調整領用量並逐週核對',target='標準化耗用回到核定目標',cause_category='集中領用',evidence='領用單 001')

def test_mean3_freeze_normalized_and_missing(client):
    d=listing(client); assert d['baseline']=='mean3'
    t,k,x=item(client)
    assert x['display_status']=='待填原因' and x['data']['baseline_periods']==['2026-08','2026-07','2026-06']
    assert D(x['data']['baseline'])==10500
    assert D(x['data']['standard30'])==D(31500)/1200*30
    assert D(x['data']['standard30_baseline'])==D(31500)/3000*30
    frozen=deepcopy(x['data']); item(client,baseline='previous')
    assert item(client)[2]['data']==frozen
    r=listing(client,'2026-09','year'); assert r['data_gaps'][0]['missing']==['2025-09']
    assert len(client.get(f"/api/monthly-cases/{t['id']}/trend",params=dict(item_key=k,month='2026-09')).json()['points'])==12

def test_roles_and_explicit_follow_choice(client):
    sign(client,'view099');act(client,'reason',expected=403,**PLAN)
    sign(client,'front099');act(client,'approve',expected=403)
    act(client,'reason',expected=422,cause_category='集中領用')
    act(client,'reason',expected=422,follow_up='yes',cause_category='集中領用')
    act(client,'reason',expected=422,follow_up='no',cause_category='最低訂購量')
    t,k,x=item(client); assert x['display_status']=='待填原因'
    sign(client,'other099');r=client.post('/api/monthly-cases/'+str(t['id']),json=dict(action='reason',keys=[k],revision=t['revision'],note='越權'));assert r.status_code==403
    sign(client,'front099')

def test_reason_review_return_improvement_and_closure(client):
    act(client,'reason',**PLAN);assert item(client)[2]['display_status']=='待審核'
    sign(client,'acct099');act(client,'reject');assert item(client)[2]['display_status']=='待填原因'
    sign(client,'front099');act(client,'reason',**PLAN)
    sign(client,'acct099');act(client,'approve',expected=422);act(client,'approve_follow');assert item(client)[2]['display_status']=='改善中'
    sign(client,'front099');act(client,'progress',evidence='核對單 002');assert item(client)[2]['progress_note']=='測試查核說明'
    act(client,'submit',result_period='2026-10',evidence='');assert item(client)[2]['display_status']=='待確認結案'
    assert item(client)[2]['result_evidence']==''
    sign(client,'acct099');act(client,'reject');assert item(client)[2]['return_stage']=='result'
    sign(client,'front099');act(client,'reason',expected=422,**PLAN);act(client,'submit',result_period='2026-10',evidence='')
    sign(client,'acct099');act(client,'finish'); x=item(client)[2]
    assert x['display_status']=='已完成' and x['verified_data']['used']=='42000.00'
    assert len(x['history'])>=10 and x['data']['used']=='31500.00'
    act(client,'reopen');assert item(client)[2]['display_status']=='改善中'

def test_link_repeated_anomaly_and_scope(client):
    sign(client,'front099'); t,k,x=item(client,month='2026-10')
    assert x['link_candidates']
    original=x['data']; target=x['link_candidates'][0]
    act(client,'link',month='2026-10',target_case_id=target['case_id'],target_key=target['key'])
    t,k,x=item(client,month='2026-10');assert x['linked_to']['case_id']==target['case_id'] and x['display_status']=='改善中'
    assert x['data']==original
    act(client,'reason',month='2026-10',expected=422,**PLAN)
    sign(client,'acct099')

def test_non_follow_reason_and_review_no_automatic_close(client):
    sign(client,'front099');act(client,'reason',kind='idle',code='I',follow_up='no',cause_category='安全備貨',evidence='備貨核准單')
    assert item(client,'idle',code='I')[2]['display_status']=='待審核'
    sign(client,'acct099');act(client,'approve',kind='idle',code='I');assert item(client,'idle',code='I')[2]['completion_kind']=='免追蹤'

def test_inventory_and_purchase_rules_unchanged(client):
    data=listing(client);kinds={t['kind'] for t in data['cases']};assert {'usage','purchase','high','low','idle'}<=kinds
    assert item(client,'low',code='L')[2]['data']['threshold']['low_days']=='14'
    assert item(client,'high',code='L')[2]['data']['threshold']['high_days']=='90'

def test_stale_revision_and_source_correction_preserve_original(client):
    sign(client,'front099');t,k,x=item(client,'purchase',code='U'); frozen=x['data']
    body=dict(revision=t['revision']-1,keys=[k],action='reason',note='說明',follow_up='no',cause_category='供應商調價')
    assert client.post('/api/monthly-cases/'+str(t['id']),json=body).status_code==409
    with db.connect() as c:
        c.execute("UPDATE lines SET body=jsonb_set(body,'{in_a}','\"1500\"') WHERE batch_id IN(SELECT id FROM batches WHERE hospital='SM' AND period='2026-09') AND body->>'code'='U'")
    x=item(client,'purchase',code='U')[2]
    assert x['data']==frozen and x['source_changed'] and x['display_status']=='待填原因'
    sign(client,'acct099')

def test_stats_export_and_missing_history(client):
    d=client.get('/api/monthly-statistics',params=dict(hospitals='SM',year=2026)).json();s=d['summary']
    assert s['repeated']>0 and s['due']>=1 and s['overdue']>=1
    assert s['review']['returned']>=2 and s['completion_rate'] is not None
    assert d['quarters']['3']['unknown_history']>0
    for path,params in [('/api/monthly-statistics-export',dict(hospitals='SM',year=2026)),('/api/monthly-cases-export',dict(hospitals='SM',month='2026-10',view='follow'))]:
        r=client.get(path,params=params);assert r.status_code==200,r.text
        w=load_workbook(BytesIO(r.content));assert len(w.sheetnames)>1
    assert client.get('/api/monthly-statistics',params=dict(hospitals='DE',year=2026)).status_code==403

@pytest.mark.parametrize('legacy,follow,done,expected', [('改善中',False,False,'待填原因'),('待會計審核',False,False,'待審核'),('合理原因已核准',True,False,'改善中'),('合理原因已核准',False,False,'已完成'),('已驗證改善',True,True,'已完成')])
def test_legacy_mapping(legacy,follow,done,expected):
    assert wf.status(dict(status=legacy,follow_up=follow,follow_up_completed=done))==expected

def test_quarter_cutoff_due_rates_and_no_duplicate_improvement():
    data=dict(code='X',name='測試',ending='100',used='100',difference='10')
    x=dict(status='改善中',workflow_status='改善中',data=data,follow_up=True,deadline='2026-09-20',cause_approved=True,approved_cause=dict(category='需求下降'))
    wf.record(x,'approve_follow',dict(name='會計'),'2026-08-01T10:00:00+08:00')
    x.update(status='已完成',workflow_status='已完成',completed_at='2026-10-01T10:00:00+08:00');wf.record(x,'finish',dict(name='會計'),'2026-10-01T10:00:00+08:00')
    rows=[dict(id=1,hospital='SM',kind='high',period='2026-08',body=dict(items={'X|個':x}))]
    s=wf.stats(rows,'2026-07-01','2026-09-30');assert s['due']==1 and s['completed']==0 and s['overdue']==1
    s=wf.stats(rows,'2026-01-01','2026-12-31');assert s['completed']==1 and s['ontime']==0

def test_js_new_role_state_actions():
    source=(ROOT/'app/static/monthly-workflow.js').read_text(encoding='utf-8')
    # 僅取純權限函式，在瀏覽器之外核對 UI 與後端責任。
    start=source.index('function mwImprovement(');end=source.index('renderTracking=async function')
    script=source[start:end]+"\nprocess.stdout.write(JSON.stringify([mwAllowed('frontline',{workflow_status:'改善中',follow_up:true,cause_approved:true}),mwAllowed('accountant',{workflow_status:'待審核',follow_up:true}),mwAllowed('frontline',{workflow_status:'待確認結案'}),mwAllowed('viewer',{workflow_status:'待審核'})]));"
    assert json.loads(subprocess.check_output(['node','-e',script],text=True))==[['progress','submit'],['approve_follow','reject'],[],[]]

def test_quarter_reason_parts_reconcile_with_annual_confirmations():
    data=dict(code='X',name='測試',ending='100',used='100',difference='10')
    x=dict(status='已完成',workflow_status='已完成',data=data,follow_up=False,approved_cause=dict(category='集中領用'),completed_at='2026-10-01T10:00:00+08:00')
    wf.record(x,'approve',dict(name='會計'),'2026-10-01T10:00:00+08:00')
    rows=[dict(id=1,hospital='SM',kind='usage',period='2026-09',body=dict(items={'X|組':x}))]
    s=wf.stats(rows,'2026-01-01','2026-12-31');assert s['causes'][0]['quarters']['3']==1
    assert sum(s['causes'][0]['quarters'].values())==s['causes'][0]['count']
    assert wf.stats(rows,'2026-07-01','2026-09-30')['unknown_history']==1

def test_batch_validation_leaves_all_items_unchanged(monkeypatch):
    monkeypatch.setattr(mc,'scoped',lambda *a,**kw:None)
    a=dict(status='待審核',workflow_status='待審核',present=True,data=dict(valid=True),submission='reasonable',cause_category='集中領用',follow_up=False,note='原因')
    b={**a,'present':False};t=dict(id=1,hospital='SM',kind='usage',revision=1,body=dict(items={'a':a,'b':b}));before=deepcopy(t)
    with pytest.raises(HTTPException):mc.act(None,t,dict(keys=['a','b'],action='approve',revision=1,note='審核'),dict(role='accountant',name='會計',id=1))
    assert t==before

def test_export_matches_keyword_and_record_filters(client):
    sign(client,'acct099')
    r=client.get('/api/monthly-cases-export',params=dict(hospitals='SM',month='2026-10',view='all',keyword='不存在的品項'))
    w=load_workbook(BytesIO(r.content));assert w['案件明細'].max_row==1
    r=client.get('/api/monthly-cases-export',params=dict(hospitals='SM',month='2026-10',view='records',date_from='2099-01-01'))
    w=load_workbook(BytesIO(r.content));assert w['操作紀錄'].max_row==1

def test_reviewer_can_correct_wrong_link(client):
    sign(client,'acct099');act(client,'unlink',month='2026-10')
    x=item(client,month='2026-10')[2];assert 'linked_to' not in x and x['display_status']=='待填原因'
    sign(client,'front099');target=x['link_candidates'][0]
    act(client,'link',month='2026-10',target_case_id=target['case_id'],target_key=target['key'])

def test_v098_pending_case_gets_current_migration_evidence(client):
    sign(client,'acct099');t,k,x=item(client,'purchase',code='U');raw=deepcopy(x)
    for name in ('workflow_status','history','source_data','latest_data','display_status','link_candidates','comparisons'):
        raw.pop(name,None)
    raw.update(status='改善中',note='舊案原始說明',follow_up=False,needs_reconfirm=False)
    with db.connect() as c:
        row=c.execute('SELECT body FROM monthly_cases WHERE id=%s',(t['id'],)).fetchone();body=row['body'];body['items'][k]=raw
        c.execute('UPDATE monthly_cases SET body=%s WHERE id=%s',(Jsonb(body),t['id']))
    x=item(client,'purchase',code='U')[2]
    assert x['display_status']=='待填原因' and x['note']=='舊案原始說明'
    assert x['history'][0]['action']=='migrated'
    assert wf.as_of(x,datetime.now(wf.TZ).date().isoformat()) is not None
    assert wf.as_of(x,'2026-09-30') is None
