"""存貨Excel真實資料列、退回狀態整併、本期統計及匯出一致性。"""
from io import BytesIO
from copy import deepcopy
import json
from decimal import Decimal as D
from openpyxl import load_workbook
import pytest
from test_workflow import client,sign,listing,act,item,PLAN
from app import monthly_workflow as wf,monthly_cases as mc

@pytest.mark.parametrize('kind',['high','low','idle','all'])
def test_stock_export_with_populated_rows(client,kind):
 sign(client,'admin099')
 r=client.get('/api/export/stock-grid',params=dict(hospitals='SM',month='2026-09',kind=kind,orders=json.dumps({k:dict(key='code',descending=False) for k in ['high','low','idle']})))
 assert r.status_code==200,r.text[:300]
 assert 'spreadsheetml' in r.headers['content-type']
 w=load_workbook(BytesIO(r.content));assert len(w.sheetnames)==(3 if kind=='all' else 1)
 assert all(s.max_row>1 for s in w)
 for sheet in w:
  assert sheet['D2'].number_format=='#,##0'
  assert sheet['E2'].number_format=='#,##0.00'


def test_legacy_return_states_map_without_losing_notes():
 for stage,want in [('reason','待填原因'),('result','改善中')]:
  x=dict(workflow_status='待補充',return_stage=stage,review_note='補充查核紀錄',history=[{'test':'保留'}]);original=deepcopy(x)
  y=wf.normalize(x);assert wf.status(y)==want and y['review_note']==x['review_note'] and y['history']==x['history']
  assert x==original
 assert '待補充' not in wf.STATES


def test_new_rejections_use_existing_states_and_grouped_exports(client):
 sign(client,'front099');act(client,'reason',**PLAN)
 sign(client,'acct099');act(client,'reject')
 t,k,x=item(client);assert x['display_status']=='待填原因' and x['review_note']
 assert mc.in_view(t,x,'follow','2026-09') and mc.matches_state(x,'改善中','follow')
 sign(client,'front099');act(client,'reason',**PLAN);act(client,'submit',result_period='2026-09',evidence='成果單112')
 t,k,x=item(client);assert mc.matches_state(x,'待審核','follow')
 for state,n in [('待審核',2),('改善中',1)]:
  r=client.get('/api/monthly-cases-export',params=dict(hospitals='SM',month='2026-09',view='follow',domain='usage',state=state));assert r.status_code==200
  assert load_workbook(BytesIO(r.content))['案件明細'].max_row==n
 sign(client,'acct099');act(client,'reject')
 t,k,x=item(client);assert x['display_status']=='改善中' and x['return_stage']=='result'
 sign(client,'front099');act(client,'submit',result_period='2026-09',evidence='補正單112')
 sign(client,'acct099');act(client,'finish');assert item(client)[2]['display_status']=='已完成'


def test_month_statistics_and_excel_use_selected_analysis_month(client):
 sign(client,'admin099');listing(client,'2026-09');listing(client,'2026-10')
 q=dict(hospitals='SM',year=2026,month='2026-09',domain='usage')
 r=client.get('/api/monthly-statistics',params=q);assert r.status_code==200,r.text
 d=r.json();assert d['month']=='2026-09' and d['summary']['start']=='2026-09-01' and d['summary']['end']=='2026-09-30'
 assert d['summary']['new_count']==1
 assert all(x['period']<='2026-09' for x in d['summary']['details'])
 q['month']='2026-10';d2=client.get('/api/monthly-statistics',params=q).json();assert d2['summary']['start']=='2026-10-01' and d2['summary']['new_count']==1
 r=client.get('/api/monthly-statistics-export',params={**q,'month':'2026-09'});assert r.status_code==200
 w=load_workbook(BytesIO(r.content));scope=dict(w['統計口徑'].values);assert scope['期間']=='本期 2026-09' and scope['起始日']=='2026-09-01' and scope['截止日']=='2026-09-30'
 r=client.get('/api/monthly-statistics',params=dict(hospitals='SM',year=2026,quarter=3));assert r.status_code==200 and r.json()['summary']['start']=='2026-07-01'
 assert client.get('/api/monthly-statistics',params={**q,'month':'2026-13'}).status_code==422
