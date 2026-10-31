"""管理總覽與管理報表摘要共用相同原精度數據，格式只作用於顯示。"""
from .dashboard import build_dashboard
from openpyxl.chart import LineChart,BarChart,Reference


def extend(wb,reports,overview,sheet):
 dash=overview.get('dashboard') or build_dashboard(reports,overview)
 scope='、'.join(dash['hospitals']);base='、'.join(dash['baseline_periods'])
 sheet(wb,'總覽四項指標',['指標','本期數值','基準數值','差異','增減%','比較用本期','共同院別','狀態及範圍'],[[c['title'],c['value'],c['baseline'],c['change']['amount'],c['change']['pct'] if c['change']['pct'] is not None else '基準為零' if c['baseline']==0 else '基準為負值' if c['baseline'] is not None and c['baseline']<0 else '缺少基準資料',c['comparison_current'],'、'.join(c['comparison_hospitals']),c['reason']+'；實際涵蓋'+str(len(c['covered']))+'/'+str(len(c['requested']))+'院；缺少'+('、'.join(c['missing']) or '無')+'；'+c['comparison_note']] for c in dash['cards']])
 sheet(wb,'總覽管理摘要',['分析年月','醫院範圍','比較月份','管理摘要'],[[dash['period'],scope,base,text] for text in dash['paragraphs']])
 def rows(block):return [[x['category'],x['current'],x['baseline'],x['difference'],x['pct']] for x in block['rows']]+[[x['category'],'—','—','—',x['reason']] for x in block['excluded']]
 sheet(wb,'耗用差異前5名',['中分類','本期耗用','基準期耗用','增減金額','增減%'],rows(dash['ranks']['consumption']))
 sheet(wb,'庫存差異前5名',['中分類','本期末庫存','基準期末庫存','增減金額','增減%'],rows(dash['ranks']['inventory']))
 ws=sheet(wb,'總覽趨勢',['月份','每月耗用總額','每床標準月均耗用','每月底庫存金額','每月存貨週轉天數','缺漏院別','範圍說明'],[[x['period'],x['used'],x['standard30_per_bed'],x['ending'],x['days'],'、'.join(x['missing']),x['scope_reason']] for x in dash['trend'][-12:]])
 for col,title,anchor,typ in [(2,'每月耗用總額','I2',BarChart),(3,'每床標準月均耗用','I18',LineChart),(4,'每月底庫存金額','I34',BarChart),(5,'每月存貨週轉天數','I50',LineChart)]:
  ch=typ();ch.title=title;ch.display_blanks='gap';ch.add_data(Reference(ws,min_col=col,min_row=1,max_row=ws.max_row),titles_from_data=True);ch.set_categories(Reference(ws,min_col=1,min_row=2,max_row=ws.max_row));ch.y_axis.numFmt='0.0' if col==5 else '#,##0';ws.add_chart(ch,anchor)
 # 管理報表沿用先前要求：兩組「增減金額由大到小」的前5名。
 summary=wb['管理摘要'];wb.remove(summary)
 records=[]
 for kind,title in [('consumption','耗用中分類差異前5名'),('inventory','期末庫存中分類差異前5名')]:
  block=dash['report_ranks'][kind]
  records += [[dash['period'],scope,base,title,x['rank'],x['category'],x['current'],x['baseline'],x['difference'],x['pct'],'依增減金額由大到小'] for x in block['rows']]
  records += [[dash['period'],scope,base,title,None,x['category'],None,None,None,None,x['reason']] for x in block['excluded']]
  if not block['rows'] and not block['excluded']:records.append([dash['period'],scope,base,title,None,'—',None,None,None,None,'尚無可比較分類資料'])
 sheet(wb,'管理摘要',['分析年月','醫院範圍','比較月份','分析項目','名次','中分類','本月金額','基準金額','增減金額','增減%','說明'],records)
 sheet(wb,'追蹤入口口徑',['項目','數值'],[['待確認提醒',dash['entries']['pending_count']],['逾期改善',dash['entries']['overdue_count']],['狀態截至日',dash['entries']['as_of']],['統計口徑',dash['entries']['note']]])
 return dash


def format_dashboard(wb):
 for ws in wb:
  headers=[str(c.value or '') for c in ws[1]];pairs=len(headers)==2 and headers[0] in ('項目','類型')
  for row in ws.iter_rows(min_row=2):
   for i,c in enumerate(row):
    label=str(row[0].value or '') if pairs and i==1 else headers[i] if i<len(headers) else ''
    if '%' in label or any(x in label for x in ('增減%','完成率','達成率','偏差率')) or ('天數' in label and '日曆天數' not in label):c.number_format='0.0'
 for name in ['總覽四項指標','總覽趨勢','耗用差異前5名','庫存差異前5名','管理摘要']:
  if name not in wb:continue
  ws=wb[name]
  for row in ws.iter_rows(min_row=2):
   for i,c in enumerate(row):
    label=str(ws.cell(1,i+1).value or '')
    if label=='增減%' or label=='每月存貨週轉天數':c.number_format='0.0'
    elif name=='總覽四項指標' and i in (1,2,3,5):c.number_format='0.0' if row[0].value=='存貨週轉天數' else '#,##0'
