"""Excel明細、圖表及來源證據輸出。"""
from decimal import Decimal
from .engine import calc
from openpyxl.chart import BarChart,LineChart,Reference
from openpyxl.chart.label import DataLabelList


def extend_workbook(wb,d,sheet,ai=None):
 summary=wb['彙總'];summary.append(['比較期',','.join(d['baseline'])]);summary.append(['採購比較基準',d['purchase_baseline']]);summary.append(['完整門檻設定',str(d['rules'])]);summary.append(['高庫存占比',d['total']['high_inventory_pct']]);summary.append(['異常品項數',d['total']['anomaly_items']]);summary.append(['資料範圍','僅啟用版本；搜尋不影響整院匯出；缺資料不當作0'])
 sheet(wb,'跨期耗用差異',['品號','品名','分類','當月耗用','比較期耗用','增減金額','增減率%','當月數量','比較期數量','數量差異','當月單位耗用成本','比較期單位耗用成本','與上月數量因素','與上月單位成本因素','拆解原因／確認理由'],[[x['code'],x['name'],x['category'],x['used'],x['baseline_used'],x['change']['amount'],x['change']['pct'],x['qty'],x['baseline_qty'],x['qty_change']['amount'],x['unit_cost'],x['baseline_unit_cost'],x['factor'].get('quantity'),x['factor'].get('unit_cost'),x['factor'].get('reason') or (x.get('factor_confirmation') or {}).get('reason','符合可比條件')] for x in d['items']])
 sheet(wb,'採購價差',['品號','品名','單位','進貨量','本月含稅單價','上次有效進貨月','上次單價','近3次加權單價','本次比較月份','比較基準單價','價差','增減率%','價差參考金額非損益','觸發門檻','資料狀態'],[[x['code'],x['name'],x['unit'],Decimal(x['in_q']),x['purchase'],x['last_purchase_period'],x['last_purchase'],x['weighted3_purchase'],','.join(x['purchase_periods']),x['purchase_basis'],x['purchase_delta']['amount'],x['purchase_delta']['pct'],x['purchase_impact'],'是' if x['purchase_alert'] else '否',x['purchase_reason'] or ('有效' if x['purchase_basis_complete'] else '比較資料不足')] for x in d['items']])
 sheet(wb,'分類比較',['分類','耗用金額','耗用占比%','增減金額','增減率%','期末存貨','庫存占比%','平均存貨','週轉率','週轉天數','主要差異品號'],[[x['category'],x['used'],x['usage_share'],x['change']['amount'],x['change']['pct'],x['ending'],x['inventory_share'],x['average'],x['turnover'],x['days'],','.join(x['top_items'])] for x in d['categories']])
 tr=sheet(wb,'月趨勢與床數',['年月','耗用','期末存貨','週轉天數','平均住院床數自動對應範圍','當月每床月耗用自動對應範圍','計算範圍'],[[x['period'],x['used'],x['ending'],x['days'],x['average_beds'],x['per_bed'],x['care_scope_name']] for x in d['trend']])
 for col,title,pos in [(2,'耗用趨勢','I2'),(3,'存貨趨勢','I18'),(4,'週轉天數','I34'),(6,'每床月耗用','I50')]:
  if tr.max_row>1:
   chart=LineChart();chart.title=title;chart.display_blanks='gap';chart.add_data(Reference(tr,min_col=col,min_row=1,max_row=tr.max_row),titles_from_data=True);chart.set_categories(Reference(tr,min_col=1,min_row=2,max_row=tr.max_row));tr.add_chart(chart,pos)
 abc=sorted(d['items'],key=lambda x:x['used'],reverse=True)
 ws=sheet(wb,'ABC與Top10',['品號','品名','耗用金額','累積占比%','ABC','期末存貨'],[[x['code'],x['name'],x['used'],x['cumulative_pct'],x['abc'],x['ending']] for x in abc])
 if abc:
  chart=BarChart();chart.title='ABC帕累托（完整品項）';chart.add_data(Reference(ws,min_col=3,min_row=1,max_row=ws.max_row),titles_from_data=True);chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=ws.max_row))
  line=LineChart();line.y_axis.axId=200;line.y_axis.title='累積占比%';line.y_axis.scaling.min=0;line.y_axis.scaling.max=100;line.y_axis.crosses='max';line.add_data(Reference(ws,min_col=4,min_row=1,max_row=ws.max_row),titles_from_data=True);chart+=line;ws.add_chart(chart,'H2')
  for val,col,title,pos in [('used',3,'耗用Top10','H20'),('ending',6,'庫存Top10','H38')]:
   top=sheet(wb,'耗用Top10' if val=='used' else '庫存Top10',['品名','金額'],[[x['name'],x[val]] for x in sorted(abc,key=lambda x:x[val],reverse=True)[:10]])
   b=BarChart();b.type='bar';b.title=title;b.add_data(Reference(top,min_col=2,min_row=1,max_row=top.max_row),titles_from_data=True);b.set_categories(Reference(top,min_col=1,min_row=2,max_row=top.max_row));ws.add_chart(b,pos)
 sheet(wb,'長期未耗用檢查',['品號','品名','期末存貨','已確認連續無耗用天數','達長期門檻且有庫存','判斷依據','持續未耗用提醒','提醒程度'],[[x['code'],x['name'],x['ending'],x['idle_observed_days'],'是' if x['idle_confirmed'] else '否',x['idle_reason'],'是' if x.get('idle_persistent') else '否',x.get('idle_level','')] for x in d['items']])
 ss=sheet(wb,'三層數據摘要',['層級','品號或分類','文字','比較期','規則版本','摘要規則','來源連結'],[[x['level'],x['code'],x['text'],','.join(x['baseline']),x['rule_version'],x['rule'],'查看來源證據'] for x in d['summaries']])
 from .excel import FIELDS
 source_ws=sheet(wb,'原始數據與計算',['批次','品號','來源品名','主檔品名','來源單位','年月','檔名','工作表','原始列','版本','有效',*FIELDS.values(),'耗用金額','耗用數量'],[[x['batch_id'],x['code'],x.get('source_name',x['name']),x['name'],x.get('source_unit',x['unit']),x['period'],x.get('filename'),x['sheet'],x['row'],x['version'],x['valid'],*[Decimal(str(x[k])) if x.get(k) is not None else None for k in FIELDS],*((calc(x)[k] for k in ('used','qty')) if x['valid'] else [None,None])] for x in d['source_rows']])
 source_index={(x['batch_id'],x['row']):i for i,x in enumerate(d['source_rows'],2)}
 ev=sheet(wb,'摘要來源證據',['摘要編號','層級','對象','來源品號','品名','院別','年月','資料版本','分類版本','工作表','原始列','批次','公式','原始數據連結'],[])
 for idx,x in enumerate(d['summaries'],start=2):
  first=ev.max_row+1
  for source in x['sources']:
   ev.append([idx-1,x['level'],x['code'],source['code'],source['name'],x['hospital'],source['period'],source['version'],source['category_version'],source['sheet'],source['row'],source['batch_id'],x['formula'],'查看原始數值'])
   ev.cell(ev.max_row,14).hyperlink=f"#'原始數據與計算'!A{source_index.get((source['batch_id'],source['row']),1)}";ev.cell(ev.max_row,14).style='Hyperlink'
  ss.cell(idx,7).hyperlink=f"#'摘要來源證據'!A{first}";ss.cell(idx,7).style='Hyperlink'
 for row in ev:
  for cell in row:
   if isinstance(cell.value,str) and cell.value.startswith(('=','+','-','@')):cell.value="'"+cell.value
 sheet(wb,'改善管理摘要',['對象','追蹤ID','修訂','摘要'],[[x['code'],x['review_id'],x['revision'],x['text']] for x in d['improvement_summary']])
 sheet(wb,'門檻與口徑',['項目','設定'],[['門檻版本',d['rule_version']],['生效月份',d['rules'].get('effective_from','舊版規則')],['門檻內容',d['rules']],['資料不足與口徑',d['notes']],['床數比較',d['bed_comparison']]])
 sheet(wb,'改善後續觀察',['追蹤ID','登錄月份','觀察月份','對象','登錄月耗用','觀察月耗用','差額','增減%','措施','狀態','解讀限制'],[[x['review_id'],x['origin_period'],x['period'],x['code'],x['before'],x['after'],x['change']['amount'],x['change']['pct'],x['measure'],x['status'],x['note']] for x in d['improvement_followup']])
 if ai:ai={**ai,'output':__import__('app.formatting',fromlist=['money_in_text']).money_in_text(ai['output'])};sheet(wb,'選定AI摘要',['項目','內容'],[['摘要ID',ai['id']],['模型',ai['model']],['產生時間',str(ai['created_at'])],['覆核','已覆核；匯出時來源未過期'],['摘要',ai['output']],['比較期',ai['input'].get('baseline')]])
 for ws in wb:
  for row in ws:
   for cell in row:
    if isinstance(cell.value,(Decimal,float)):cell.number_format='#,##0.00####'
  ws.print_options.horizontalCentered=True;ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
