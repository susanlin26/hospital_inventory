"""Tracking-only routes. No new financial rules or database schema."""
from datetime import date
from decimal import Decimal as D
from io import BytesIO
from fastapi import Depends,HTTPException,Response
from openpyxl import Workbook
from openpyxl.styles import Alignment
from . import auth,db,excel,tracking,tracking_desk as desk

def register(app,prepare,jsonable,sheet,report_func):
 def scope(hospitals,month,baseline,u):
  hs=list(dict.fromkeys(hospitals.split(',')))
  if not hs or len(hs)>50 or baseline not in ('previous','year'):raise HTTPException(422,'查詢條件不正確')
  for h in hs:auth.authorize(u,h)
  return hs,excel.period(month)
 def header(c,hs,p,base):
  rows=c.execute("SELECT hospital,identity,rule,level,active,case_id,link_origin,fingerprint,confirmed_fingerprint,confirmation IS NOT NULL confirmation,jsonb_build_object('code',snapshot->>'code','name',snapshot->>'name','unit',snapshot->>'unit') snapshot FROM tracking_observations WHERE hospital=ANY(%s) AND period=%s AND scope='__all__' AND basis IN (%s,'fixed')",(hs,p,base)).fetchall()
  normal=[o for o in rows if o['rule'] not in ('high_coverage','長期未耗用') and o['active'] and not desk.data_issue(o) and o['level']=='品項' and desk.pending(o)]
  counts=c.execute("SELECT count(*) n,count(*) FILTER(WHERE body->>'due_date'<(now() AT TIME ZONE 'Asia/Taipei')::date::text) overdue FROM tracking_cases WHERE hospital=ANY(%s) AND NOT(status=ANY(%s))",(hs,list(tracking.CLOSED))).fetchone()
  missing=c.execute("SELECT DISTINCT hospital FROM batches b WHERE b.hospital=ANY(%s) AND b.period=%s AND b.active AND b.deleted_at IS NULL AND NOT EXISTS(SELECT 1 FROM tracking_scans s WHERE s.hospital=b.hospital AND s.period=b.period AND s.scope='__all__' AND s.basis=%s)",(hs,p,base)).fetchall()
  return dict(pending_ready=not missing,pending=len({desk.product_key(o) for o in normal}),cases=counts['n'],overdue=counts['overdue'],global_alerts=sum(o['active'] and not desk.data_issue(o) and o['level']!='品項' for o in rows),data_issues=sum(o['active'] and desk.data_issue(o) for o in rows))
 def listing(c,hs,p,base,tab,kind,category,processing,q,rule,issues,unit,status,overdue,record_kind,date_from,date_to,item_code=''):
  result={'tab':tab,'period':p,'baseline':base,'hospitals':hs}
  if tab=='confirm':
   prepared=prepare(c,hs,p,base,'__all__',True,summary_only=True);reports=prepared.pop('_desk_reports')
   from .case_problems import decorate
   rows=decorate(c,desk.observations(c,hs,p,base));matched=desk.filtered(rows,kind,category,processing,q,rule,issues);matched=[o for o in matched if not item_code or o['snapshot'].get('code')==item_code]
   result.update(observations=rows,matched_ids=[o['id'] for o in matched],groups=desk.groups(matched),audit=desk.audit(reports,base),baseline_periods=prepared['baseline_periods'],history_pending=prepared['history_pending'],history_jobs=prepared['history_jobs'],categories=prepared['categories'])
  elif tab=='cases':result['cases']=desk.filter_cases(desk.case_list(c,hs),kind,unit,status,overdue,q)
  elif tab=='closed':
   if record_kind not in ('reasonable','completed','legacy'):raise HTTPException(422,'紀錄類別不正確')
   result['records'],result['date_from'],result['date_to']=desk.records(c,hs,record_kind,date_from,date_to,q);result['record_kind']=record_kind
  else:raise HTTPException(422,'分頁不正確')
  result['counts']=header(c,hs,p,base);return result

 @app.get('/api/tracking/desk')
 def list_desk(hospitals:str,month:str,baseline:str='previous',tab:str='confirm',kind:str='all',category:str='',processing:str='pending',q:str='',rule:str='',issues:bool=False,unit:str='',status:str='',overdue:bool=False,record_kind:str='reasonable',date_from:str='',date_to:str='',item_code:str='',u=Depends(auth.current)):
  hs,p=scope(hospitals,month,baseline,u)
  with db.connect() as c:return jsonable(listing(c,hs,p,baseline,tab,kind,category,processing,q,rule,issues,unit,status,overdue,record_kind,date_from,date_to,item_code))

 @app.get('/api/tracking/desk/observations/{oid}')
 def observation(oid:int,u=Depends(auth.current)):
  with db.connect() as c:
   o=c.execute('SELECT * FROM tracking_observations WHERE id=%s',(oid,)).fetchone()
   if not o:raise HTTPException(404,'提醒不存在')
   auth.authorize(u,o['hospital']);row=desk.compact(o);row['snapshot']=o['snapshot'];row['confirmation']=o['confirmation']
   row['related_cases']=[{k:t[k] for k in ('id','title','status')} for t in c.execute("SELECT DISTINCT t.id,t.title,t.status FROM tracking_cases t JOIN tracking_observations a ON a.case_id=t.id WHERE t.hospital=%s AND a.identity=%s AND a.rule=%s AND a.scope=%s AND a.basis=%s",(o['hospital'],o['identity'],o['rule'],o['scope'],o['basis']))]
   from .case_problems import decorate
   row=decorate(c,[row])[0];row['related_cases']=row['candidate_cases']
   return jsonable(row)

 @app.get('/api/tracking/desk/legacy/{legacy_id}')
 def legacy_case(legacy_id:int,u=Depends(auth.current)):
  with db.connect() as c:
   t=c.execute('SELECT id,hospital FROM tracking_cases WHERE legacy_id=%s',(legacy_id,)).fetchone()
   if not t:raise HTTPException(404,'尚無此舊案的相容紀錄，請先開啟追蹤清單同步')
   auth.authorize(u,t['hospital']);return {'id':t['id']}

 @app.get('/api/tracking/desk/cases/{cid}')
 def case(cid:int,u=Depends(auth.current)):
  with db.connect() as c:
   t=c.execute('SELECT * FROM tracking_cases WHERE id=%s',(cid,)).fetchone()
   if not t:raise HTTPException(404,'案件不存在')
   auth.authorize(u,t['hospital']);from .case_problems import summary
   return jsonable({**t,**summary(c,t)})

 @app.get('/api/tracking/desk/cases/{cid}/observations')
 def case_observations(cid:int,u=Depends(auth.current)):
  with db.connect() as c:
   t=c.execute('SELECT hospital FROM tracking_cases WHERE id=%s',(cid,)).fetchone()
   if not t:raise HTTPException(404,'案件不存在')
   auth.authorize(u,t['hospital']);return jsonable([desk.compact(o) for o in c.execute('SELECT * FROM tracking_observations WHERE case_id=%s ORDER BY period,id',(cid,))])

 @app.get('/api/tracking/desk/records/{bid}')
 def record(bid:int,u=Depends(auth.current)):
  with db.connect() as c:
   b=c.execute('SELECT b.*,u.name operator_name FROM tracking_batches b LEFT JOIN users u ON u.id=b.operator WHERE b.id=%s',(bid,)).fetchone()
   if not b:raise HTTPException(404,'確認紀錄不存在')
   auth.authorize(u,b['hospital']);b['observations']=[desk.compact(o) for o in c.execute('SELECT * FROM tracking_observations WHERE id=ANY(%s) ORDER BY id',(b['body']['ids'],))];return jsonable(b)

 @app.post('/api/tracking/desk/link')
 def link(body:dict,u=Depends(auth.current)):
  if not body.get('case_id'):
   if any(not str(body.get(k,'')).strip() for k in ('title','cause','responsible_unit','measure','due_date')):raise HTTPException(422,'請填案件名稱、問題說明、負責單位、措施及預計完成日')
   date.fromisoformat(body['due_date'])
  with db.connect() as c:
   db.write_lock(c);c.execute('SELECT pg_advisory_xact_lock(7414080)');return tracking.link(c,body,u)

 def problem_context(c,t):
  from .case_problems import groups
  latest=c.execute('SELECT max(period) p FROM batches WHERE hospital=%s AND active AND deleted_at IS NULL',(t['hospital'],)).fetchone()['p'] or t['start_period']
  generated=[]
  for scope,base in {(p['scope'],p['basis'] if p['basis'] in ('previous','year') else 'previous') for p in groups(c,t).values()}:
   d=report_func(c,t['hospital'],latest,baseline=base,department=scope or None);tracking.sync(c,[d],base);generated.extend(tracking.observations(d,base))
  return generated,latest

 @app.get('/api/tracking/cases/{cid}/problems')
 def problems(cid:int,u=Depends(auth.current)):
  from .case_problems import view
  with db.connect() as c:
   t=c.execute('SELECT * FROM tracking_cases WHERE id=%s',(cid,)).fetchone()
   if not t:raise HTTPException(404,'案件不存在')
   auth.authorize(u,t['hospital']);generated,period=problem_context(c,t);t=c.execute('SELECT * FROM tracking_cases WHERE id=%s',(cid,)).fetchone()
   return jsonable({**view(c,t,generated,period),'revision':t['revision'],'status':t['status']})

 @app.post('/api/tracking/cases/{cid}/problems')
 def problem_action(cid:int,body:dict,u=Depends(auth.current)):
  from .case_problems import act
  with db.connect() as c:
   db.write_lock(c);c.execute('SELECT pg_advisory_xact_lock(7414080)');t=tracking.get_case(c,cid,u);generated,period=problem_context(c,t)
   return act(c,t,body,u,generated,period)

 @app.get('/api/tracking/problem-records/{eid}')
 def problem_record(eid:int,u=Depends(auth.current)):
  with db.connect() as c:
   e=c.execute("SELECT * FROM tracking_events WHERE id=%s AND action='案件問題整批處理'",(eid,)).fetchone()
   if not e:raise HTTPException(404,'紀錄不存在')
   auth.authorize(u,e['hospital']);return jsonable(e)

 def reasonable_payload(c,bid):
  if str(bid).startswith('problem-'):
   e=c.execute('SELECT body FROM tracking_events WHERE id=%s',(int(str(bid).split('-')[1]),)).fetchone()['body'];snapshots={}
   for change in e.get('changes',[]):
    v=e['after']['body']['problem_results'][change['key']]['history'][-1]['assessment']
    if v.get('snapshot'):snapshots[str(v.get('observation_id',''))]=v['snapshot']
   return dict(snapshots=snapshots)
  return c.execute('SELECT body FROM tracking_batches WHERE id=%s',(bid,)).fetchone()['body']

 @app.get('/api/export/tracking-desk')
 def export(hospitals:str,month:str,baseline:str='previous',tab:str='confirm',kind:str='all',category:str='',processing:str='pending',q:str='',rule:str='',issues:bool=False,unit:str='',status:str='',overdue:bool=False,record_kind:str='reasonable',date_from:str='',date_to:str='',item_code:str='',u=Depends(auth.current)):
  hs,p=scope(hospitals,month,baseline,u)
  with db.connect() as c:
   data=listing(c,hs,p,baseline,tab,kind,category,processing,q,rule,issues,unit,status,overdue,record_kind,date_from,date_to,item_code);wb=Workbook();wb.remove(wb.active)
   sheet(wb,'篩選條件',['條件','內容'],[['分頁',{'confirm':'待確認異常','cases':'改善案件','closed':'處理紀錄'}[tab]],['醫院','、'.join(hs)],['分析月份',p if tab=='confirm' else '案件不受分析月份限制' if tab=='cases' else '不適用：依處理日期'],['比較基準','上月' if baseline=='previous' else '去年同期'],['比較年月','、'.join(data.get('baseline_periods',[]))],['問題類別',{'all':'全部','usage':'耗用','stock':'存貨','both':'耗用及存貨','other':'其他提醒'}.get(kind,kind)],['中分類',category or '全部'],['處理情形',{'pending':'待確認及需重新確認','all':'全部處理情形'}.get(processing,processing)],['關鍵字',q],['指定品號',item_code],['提醒類型',rule or '全部'],['資料待補／核對',issues],['負責單位',unit],['案件狀態',status],['僅逾期',overdue],['紀錄類別',{'reasonable':'合理原因確認','completed':'改善完成','legacy':'舊案紀錄'}.get(record_kind,record_kind)],['處理日期起',data.get('date_from','')],['處理日期迄',data.get('date_to','')],['歷史提醒','背景整理中，連續月份不作完整統計' if data.get('history_pending') else '目前狀態']])
   obs=[]
   if tab=='confirm':
    obs=[o for o in data['observations'] if o['id'] in set(data['matched_ids'])]
    sheet(wb,'中分類清單',['中分類','待確認品項數','主要提醒'],[[g['category'],g['count'],'、'.join(g['labels'])] for g in data['groups']])
   elif tab=='cases':
    sheet(wb,'改善案件',['案件編號','醫院','案件名稱','問題類別','涉及品項數','負責單位','期限','狀態','逾期','最新進度'],[[t['id'],t['hospital'],t['title'],t['problem_type'],t['item_count'],t['responsible_unit'],t['due_date'],t['status'],t['overdue'],t['progress']] for t in data['cases']])
    ids=[t['id'] for t in data['cases']];obs=[desk.compact(o) for o in c.execute('SELECT * FROM tracking_observations WHERE case_id=ANY(%s) ORDER BY case_id,period,id',(ids,))]
   else:
    if record_kind=='reasonable':
     sheet(wb,'合理原因確認',['確認紀錄','醫院','確認日期','原因類別','原因摘要','涉及品項數','提醒數','確認人','目前情形'],[[b['id'],b['hospital'],b['date'],b['cause_category'],b['summary'],b['item_count'],b['alert_count'],b['operator'],b['state']] for b in data['records']])
     details=[]
     for b in data['records']:
      saved=reasonable_payload(c,b['id'])
      for oid,snap in saved.get('snapshots',{}).items():details.append([b['id'],int(oid),b['hospital'],snap.get('code'),snap.get('name'),snap.get('unit'),snap.get('label'),snap.get('periods') and '、'.join(snap['periods']),snap.get('current'),snap.get('reference'),snap.get('difference'),snap.get('value_unit'),b['summary']])
     sheet(wb,'確認當時提醒',['確認紀錄','提醒編號','醫院','品號','品名及規格','計量單位','提醒','採用月份','本期數值','基準或門檻','差異','數值單位','共同原因'],details)
    else:
     sheet(wb,'改善完成' if record_kind=='completed' else '舊案紀錄',['案件編號','原案件編號','醫院','案件名稱','負責單位','處理日期','結果摘要','紀錄狀態'],[[t['id'],t['legacy_id'],t['hospital'],t['title'],t['responsible_unit'],t['date'],t['summary'],t['status']] for t in data['records']])
     ids=[t['id'] for t in data['records']];obs=[desk.compact(o) for o in c.execute('SELECT * FROM tracking_observations WHERE case_id=ANY(%s) ORDER BY case_id,period,id',(ids,))]
   if obs or tab!='closed':
    sheet(wb,'提醒明細',['提醒編號','案件編號','醫院','分析年月','中分類','品號','品名及規格','計量單位','問題','判讀重點','本期數值','基準或門檻','差異','數值單位','處理情形','比較基準','採用月份','公式'],[[o['id'],o['case_id'],o['hospital'],o['period'],o['category'],o['snapshot'].get('code'),o['snapshot'].get('name'),o['snapshot'].get('unit'),o['snapshot']['label'],o['reading'],o['snapshot'].get('current'),o['snapshot'].get('reference'),o['snapshot'].get('difference'),o['snapshot'].get('value_unit'),o['processing'],{'previous':'上月','year':'去年同期','fixed':'固定庫存或歷史門檻'}.get(o['basis'],o['basis']),'、'.join(o['snapshot'].get('periods',[])),o['snapshot'].get('formula')] for o in obs])
   # Include selected-case fields and full provenance only during explicit export.
   case_ids=[t['id'] for t in data.get('cases',[])] if tab=='cases' else [t['id'] for t in data.get('records',[])] if tab=='closed' and record_kind!='reasonable' else []
   if case_ids:
    full=c.execute('SELECT * FROM tracking_cases WHERE id=ANY(%s) ORDER BY id',(case_ids,)).fetchall()
    sheet(wb,'案件目前資料',['案件編號','原案編號','醫院','案件名稱','起始月份','問題說明','負責單位','措施','期限','最新進度','驗證方式','觀察期間','目前結果','完成日','目前狀態'],[[t['id'],t['legacy_id'],t['hospital'],t['title'],t['start_period'],*[t['body'].get(k,'') for k in ('cause','responsible_unit','measure','due_date','progress','verification','observation_period','result','completed_date')],t['status']] for t in full])
   provenance=[]
   if obs:
    for o in c.execute('SELECT id,hospital,basis,snapshot FROM tracking_observations WHERE id=ANY(%s) ORDER BY id',([o['id'] for o in obs],)):
     for src in o['snapshot'].get('source',[]):provenance.append([o['id'],o['hospital'],o['snapshot'].get('rule_version'),src.get('period'),src.get('code'),src.get('batch_id'),src.get('version'),src.get('category_version'),src.get('sheet'),src.get('row')])
   if tab=='closed' and record_kind=='reasonable':
    for b in data['records']:
     saved=reasonable_payload(c,b['id'])
     for oid,snap in saved.get('snapshots',{}).items():
      for src in snap.get('source',[]):provenance.append([int(oid),b['hospital'],snap.get('rule_version'),src.get('period'),src.get('code'),src.get('batch_id'),src.get('version'),src.get('category_version'),src.get('sheet'),src.get('row')])
   if case_ids:
    from .case_problems import groups,number
    detail=[]
    for t in full:
     for v in groups(c,t).values():detail.append([number(t),t['title'],v['code'],v['name'],v['label'],v['status'],v['result'],v['joined_period'],'、'.join(v['periods'])])
    sheet(wb,'案件各問題結果',['案件號','案件名稱','品號','品名','問題','處理狀態','結果','加入月份','關聯月份'],detail)
   sheet(wb,'資料依據與版本',['提醒編號','醫院','規則版本','來源月份','來源品號','匯入批次','匯入版本','分類版本','工作表','列號'],provenance)
   for ws in wb:
    heads=[x.value for x in ws[1]]
    for row in ws.iter_rows(min_row=2):
     labels=dict(zip(heads,row));value_unit=labels.get('數值單位');vu=value_unit.value if value_unit else ''
     for cell in row:
      label=heads[cell.column-1]
      if label in ('本期數值','基準或門檻','差異') and cell.value is not None:
       try:cell.value=D(str(cell.value).removeprefix("'"));cell.number_format='#,##0' if vu in ('元','元／床／30日') else '#,##0.0' if vu=='天' else '#,##0.####';cell.alignment=Alignment(horizontal='center')
       except Exception:pass
   db.audit(c,u['id'],'tracking_desk_export',target=tab,detail={'hospitals':hs,'month':p,'date_from':date_from,'date_to':date_to});buf=BytesIO();wb.save(buf)
  return Response(buf.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="tracking_{tab}_{p}.xlsx"'})
