"""以欄名解析實際儲存值，不依工作表名稱或固定欄位位置。"""
from collections import Counter
from copy import deepcopy
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
import zipfile
from openpyxl import load_workbook

FIELDS = {
 'open_q':'上期結存數量','open_a':'上期結存金額','in_q':'本期進貨數量','in_a':'本期進貨金額',
 'sale_q':'本期銷貨數量','sale_a':'本期銷貨金額','count_q':'本期盤點數量','count_a':'本期盤點金額',
 'adj_q':'本期調整數量','adj_a':'本期調整金額','other_q':'其他異動數量','other_a':'其他異動金額',
 'end_q':'本期結存數量','end_a':'本期結存金額','price':'單價'}
ALIASES={'產品編號':'code','品號':'code','產品代碼':'code','產品名稱':'name','品名':'name','單位':'unit','主分類':'category','類型':'kind',**{v:k for k,v in FIELDS.items()}}
ZERO=Decimal('0')
TOLERANCE_POLICY='inventory-amount-1-v2-strict'
CLASSIFICATION_POLICY='code-master-2'

def sort_issues(issues):
 return sorted(issues,key=lambda x:{'error':0,'warning':1,'tolerance':2}.get(x.get('severity'),1))

def classify_line(line,classification,issue):
 """Code is the matching key. Master names do not change source quantity units."""
 raw=line.get('raw',{})
 line['source_name']=raw.get('name',line.get('source_name',line['name']))
 line['source_unit']=raw.get('unit',line.get('source_unit',line['unit']))
 line['source_category']=raw.get('category',line.get('source_category',''))
 line['name']=line['source_name'];line['unit']=line['source_unit']
 line['category']=line['source_category'] or '未分類'
 line['classification_version']=classification['id'] if classification else None
 candidates=classification['items'].get(line['code'],[]) if classification else []
 changes=[];status='unmatched'
 if candidates:
  names=sorted({x['name'] for x in candidates});units=sorted({x['unit'] for x in candidates});cats=sorted({x['category'] for x in candidates})
  line['master_names']=names;line['master_units']=units;line['master_categories']=cats
  show=lambda values:'、'.join('「'+x+'」' for x in values)
  if len(names)==1:
   line['name']=names[0]
   if line['source_name']!=line['name']:changes.append(f"品名不同：分類主檔{show(names)}；匯入檔「{line['source_name']}」，採主檔品名")
  elif line['source_name'] not in names:
   changes.append(f"主檔同品號有多個品名{show(names)}；匯入檔「{line['source_name']}」，品名待確認，暫保留來源品名")
  else:changes.append(f"主檔同品號有多個品名{show(names)}；匯入檔「{line['source_name']}」，沿用相符的主檔品名")
  if len(units)!=1 or line['unit'] not in units:changes.append(f"單位不同或主檔不唯一：分類主檔{show(units)}；匯入檔「{line['unit']}」，保留匯入單位，不換算數量")
  if len(cats)==1:
   line['category']=cats[0];status='matched'
   if line['source_category'] and line['source_category']!=cats[0]:changes.append(f"分類不同：分類主檔{show(cats)}；匯入檔「{line['source_category']}」，採主檔分類")
   conclusion=f"已依產品編號「{line['code']}」套用分類「{cats[0]}」"
  else:
   line['category']='分類待確認';status='ambiguous'
   changes.append(f"分類主檔同品號對應多個分類{show(cats)}；匯入檔分類「{line['source_category'] or '未提供'}」")
   conclusion='分類待確認，未任意套用其中一個分類'
  if changes:issue(line['row'],line['code'],'；'.join(changes+[conclusion]),'warning',check='classification',source_name=line['source_name'],source_unit=line['unit'],source_category=line['source_category'],master_names=names,master_units=units,master_categories=cats)
 else:
  message=f"分類主檔找不到產品編號「{line['code']}」" if classification else f"尚未匯入分類主檔，無法對應產品編號「{line['code']}」"
  issue(line['row'],line['code'],message+f"；匯入品名「{line['source_name']}」、單位「{line['unit']}」、分類「{line['source_category'] or '未提供'}」；本列隔離，請補齊主檔後重新匯入月報",'error',check='classification')
  line['valid']=False
 line['classification_status']=status

def negative_inventory(line,issue):
 for key,label,amount_key in [('open_q','期初結存數量','open_a'),('end_q','期末結存數量','end_a'),('open_a','期初結存金額','open_a'),('end_a','期末結存金額','end_a')]:
  value=line.get(key)
  if value is None:continue
  n=Decimal(value)
  if n>=0:continue
  origin=f"；匯入品名「{line['source_name']}」" if line.get('source_name')!=line['name'] else ''
  prefix=f"負庫存：工作表「{line['sheet']}」第{line['row']}列／品號「{line['code']}」／品名「{line['name']}」{origin}；"
  if key.endswith('_a'):detail=f'{label}為{n}元，低於0元共{abs(n)}元（原始未稅金額）'
  else:
   amount=line.get(amount_key);detail=f"{label}為{n}「{line['unit']}」，低於0共{abs(n)}「{line['unit']}」；對應結存金額為{amount if amount is not None else '缺漏'}元，不以數量估算金額"
  issue(line['row'],line['code'],prefix+detail+'；隔離待確認','error',check='negative_inventory',field=key,difference=str(n),shortfall=str(abs(n)),difference_unit='元' if key.endswith('_a') else line['unit'],name=line['name'])
  line['valid']=False

def refresh_identity(report,lines,classification):
 report=deepcopy(report);lines=deepcopy(lines)
 old_reasons={'分類主檔品名、單位或分類衝突，不自動套用','來源分類與主檔不同，保留來源分類並列待確認','未分類，仍納入總額','負庫存，隔離待確認'}
 report['issues']=[x for x in report.get('issues',[]) if x.get('reason') not in old_reasons and x.get('check') not in ('classification','negative_inventory')]
 def issue(row,code,reason,severity='error',**extra):report['issues'].append(dict(row=row,code=code,reason=reason,severity=severity,**extra))
 for line in lines:classify_line(line,classification,issue);negative_inventory(line,issue)
 report['issues']=sort_issues(report['issues'])
 report.update(valid_count=sum(x['valid'] for x in lines),excluded_count=sum(not x['valid'] for x in lines),classification_policy=CLASSIFICATION_POLICY,parser_version='material-1.3')
 return report,lines

def validation_status(report):
 issues=report.get('issues',[])
 if any(x.get('severity')=='error' for x in issues):return 'blocked'
 if any(x.get('severity')!='tolerance' for x in issues):return 'warning'
 return 'passed'

def apply_amount_tolerance(report):
 """Classify reconciliation precisely without changing source amounts/quantities."""
 result=deepcopy(report);amounts=[]
 for item in result.get('issues',[]):
  kind=item.get('reason');severity=item.get('severity')
  if kind not in ('進銷存金額勾稽差異','進銷存數量勾稽差異') or severity not in ('warning','tolerance','error'):continue
  try:diff=Decimal(str(item.get('difference')))
  except (InvalidOperation,ValueError,TypeError):continue
  if not diff.is_finite():continue
  if kind=='進銷存金額勾稽差異' and ZERO<abs(diff)<=Decimal('1') and severity!='error':
   item['severity']='tolerance';amounts.append(diff)
  elif diff!=0:item['severity']='error'
 result['tolerance_policy']=TOLERANCE_POLICY
 result['tolerance_summary']={'count':len(amounts),'net_difference':str(sum(amounts,ZERO)),'absolute_difference':str(sum((abs(x) for x in amounts),ZERO)),'limit':'1','basis':'原始未稅金額；僅適用逐品號進銷存金額勾稽；數量無額外容許差異'}
 return result


def period(value):
 s=str(value).strip()
 if re.fullmatch(r'\d{5}',s): s=f'{int(s[:3])+1911}-{s[3:]}'
 elif re.fullmatch(r'\d{6}',s): s=s[:4]+'-'+s[4:]
 if not re.fullmatch(r'20\d{2}-(0[1-9]|1[0-2])',s): raise ValueError('年月請使用2026-07或11507格式')
 return s

def file_identity(filename):
 """僅依完整標準檔名辨識；不可猜院別、年月或將藥品當耗材。"""
 if not isinstance(filename,str) or len(filename)>250 or '/' in filename or '\\' in filename:
  raise ValueError('檔名不正確，請使用11401_DE進銷存_耗材_RCW.xlsx格式')
 match=re.fullmatch(r'(?P<date>\d{5}|\d{6})_(?P<hospital>[A-Za-z][A-Za-z0-9]{1,11})進銷存(?:表)?_(?P<kind>耗材|藥品)(?:_(?P<department>[^_./\\]{1,60}))?\.xlsx',filename,re.IGNORECASE)
 if not match:raise ValueError('檔名格式不符，請使用11401_DE進銷存_耗材_RCW.xlsx；僅接受.xlsx')
 from .management import department
 import unicodedata
 raw_dep=unicodedata.normalize('NFKC',match['department'] or '').strip()
 if raw_dep and not re.fullmatch(r'[A-Za-z0-9\u4e00-\u9fff() -]+',raw_dep):raise ValueError('檔名科別含不允許的字元')
 dep=department(raw_dep) if raw_dep else ''
 if dep and not re.fullmatch(r'[A-Za-z0-9\u4e00-\u9fff()（） -]+',dep):raise ValueError('科別只能使用中英文字、數字、括號、空白或連字號')
 p=period(match['date']);year,month=map(int,p.split('-'))
 return {'filename':filename,'department':dep,'hospital':match['hospital'].upper(),'period':p,'roc_label':f'民國{year-1911}年{month:02d}月','category':'MATERIAL' if match['kind']=='耗材' else 'DRUG','category_label':match['kind']}

def number(v):
 if v is None or isinstance(v,bool): raise ValueError('缺漏或非數值，不可自行補零')
 try:
  n=Decimal(str(v).replace(',','').strip())
  if not n.is_finite() or abs(n)>Decimal('1e15'):raise ValueError('數值超出範圍')
  return n
 except InvalidOperation as exc: raise ValueError('不是有效數值') from exc

def clean(v): return '' if v is None else str(v).strip()

def workbook(data):
 if len(data)>15*1024*1024:raise ValueError('檔案超過15 MB上限')
 try:
  with zipfile.ZipFile(BytesIO(data)) as z:
   if sum(i.file_size for i in z.infolist())>80*1024*1024: raise ValueError('解壓後檔案超過80 MB上限')
  return load_workbook(BytesIO(data),data_only=True,read_only=False)
 except (zipfile.BadZipFile,KeyError) as exc:raise ValueError('請上傳有效的.xlsx檔案') from exc

def parse(data, classification=None):
 w=workbook(data); sheets=[]
 for s in w:
  for row in s.iter_rows(min_row=1,max_row=min(s.max_row,20)):
   mapping={ALIASES[clean(c.value)]:i for i,c in enumerate(row) if clean(c.value) in ALIASES}
   if {'code','name','unit'}.issubset(mapping) and ('open_a' in mapping):
    sheets.append((s,row[0].row,mapping));break
 if not sheets:raise ValueError('找不到包含產品編號、品名、單位與進銷存金額的工作表')
 if len(sheets)>1:raise ValueError('找到多個進銷存工作表，請分檔上傳，避免重複計入')
 s,header,mapping=sheets[0]
 if s.max_row>100000 or s.max_column>64:raise ValueError('報表超過本版支援的100,000列或64欄，請拆分或確認報表範圍')
 missing=[v for k,v in FIELDS.items() if k not in mapping]
 if missing:raise ValueError('必要欄位缺漏：'+ '、'.join(missing))
 lines=[];issues=[];blank=0;totals=[]
 def issue(row,code,reason,severity='error',**extra):
  issues.append(dict(row=row,code=code,reason=reason,severity=severity,**extra))
 for cells in s.iter_rows(min_row=header+1):
  raw=[c.value for c in cells]
  if not any(v is not None and clean(v) for v in raw):blank+=1;continue
  r=cells[0].row
  get=lambda key:raw[mapping[key]] if key in mapping and mapping[key]<len(raw) else None
  code=clean(get('code'))
  if code in ('合計','總計','總合計') or (not code and clean(get('name')) in ('合計','總計','總合計')):
   totals.append({'row':r,'values':{k:clean(get(k)) for k in FIELDS}});continue
  source={ALIASES.get(clean(s.cell(header,i+1).value),clean(s.cell(header,i+1).value)):clean(v) for i,v in enumerate(raw)}
  line=dict(code=code,name=clean(get('name')),unit=clean(get('unit')),category=clean(get('category')) or '未分類',row=r,sheet=s.title,raw=source,valid=True,classification_version=None)
  for key in ['code','name','unit']:
   if not line[key]:issue(r,code,f'{key}識別資料缺漏');line['valid']=False
  for key in FIELDS:
   try:line[key]=str(number(get(key)))
   except ValueError as e:issue(r,code,f'{FIELDS[key]}：{e}');line['valid']=False;line[key]=None
  if s.row_dimensions[r].hidden:issue(r,code,'來源列為隱藏列，保留並需檢視','warning')
  classify_line(line,classification,issue)
  negative_inventory(line,issue)
  if line['valid']:
   # 此來源的庫存勾稽：盤點、調整、其他異動均以原始帶號值加回。
   for suffix in ('q','a'):
    balance=sum(Decimal(line[k+'_'+suffix]) for k in ('open','in','count','adj','other'))-Decimal(line['sale_'+suffix])
    diff=balance-Decimal(line['end_'+suffix])
    if diff:issue(r,code,'進銷存數量勾稽差異' if suffix=='q' else '進銷存金額勾稽差異','warning',difference=str(diff),basis='期初＋進貨－銷貨＋盤點＋調整＋其他－期末')
  lines.append(line)
 counts=Counter(l['code'] for l in lines if l['code'])
 for line in lines:
  if counts[line['code']]>1:line['valid']=False;issue(line['row'],line['code'],'同院同月重複品號，隔離')
 for total in totals:
  for key,value in total['values'].items():
   if not value or key=='price':continue
   try:
    detail=sum((Decimal(l[key]) for l in lines if l[key] is not None),ZERO)
    diff=detail-number(value)
    if diff:issue(total['row'],'合計',f'{FIELDS[key]}明細與原始總計不符','error',difference=str(diff))
   except ValueError:issue(total['row'],'合計',f'{FIELDS[key]}總計非有效數字')
 result=apply_amount_tolerance(dict(lines=lines,issues=issues,sheet=s.title,header_row=header,blank_rows=blank,total_rows=totals,control_note='已檢查原始總計' if totals else '來源沒有總計列，須使用獨立對帳',parser_version='material-1.3',classification_policy=CLASSIFICATION_POLICY))
 result['issues']=sort_issues(result['issues'])
 return result

def parse_categories(data):
 w=workbook(data);items={}
 for s in w:
  rows=list(s.values)
  if not rows:continue
  head={ALIASES[clean(v)]:i for i,v in enumerate(rows[0]) if clean(v) in ALIASES}
  if not {'code','name','unit','category'}.issubset(head):continue
  for row in rows[1:]:
   obj={k:clean(row[i]) if i<len(row) else '' for k,i in head.items()}
   if not obj.get('code'):continue
   if not all(obj.get(k) for k in ('name','unit','category')):raise ValueError(f"分類主檔 {obj['code']} 缺少品名、單位或分類")
   items.setdefault(obj['code'],[]).append(obj)
  return items
 raise ValueError('分類主檔缺少產品編號、品名、主分類或單位')
