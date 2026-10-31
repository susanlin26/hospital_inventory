"""沿用既有Excel產生器，新增管理月報與追蹤表，不引入外部服務。"""
from . import engine
from .inpatient import SCOPE_NAMES


def extend_daily(wb,d,sheet):
 t=d['total'];pairs=[['當月住院人日',t['patient_days']],['未取整住院人日',t['patient_days_raw']],['當月日曆天數',t['calendar_days']],['照護範圍',t['care_scope_name']],['每住院人日耗用',t['per_patient_day']],['標準30日每床耗用',t['standard30_per_bed']],['當月每床月耗用自動對應範圍',t['monthly_per_bed']],['人日計算狀態',t['daily_status']],['前三月耗用合計',d['daily_baseline3']['used']],['前三月住院人日合計',d['daily_baseline3']['patient_days']],['前三月每人日耗用基準',d['daily_baseline3']['per_patient_day']],['前三月基準偏差率',t['daily_change3']['pct']],['資料口徑',t['scope_note']]]
 for pair in pairs:wb['彙總'].append([pair[0],str(pair[1]) if isinstance(pair[1],dict) else float(pair[1]) if hasattr(pair[1],'as_tuple') else pair[1]])
 sheet(wb,'住院人日明細',['醫院','年月','床別','原始輸入值（來源單位見狀態）','日曆天數','未取整人日','當月住院人日','來源','狀態'],[[x['hospital'],x['period'],x['department'],x.get('value'),x['calendar_days'],x['patient_days_raw'],x['patient_days'],x['source'],x['status']] for x in d['census']['records']]+[[d['hospital'],d['period'],'已提供適用範圍合計',t['average_beds'],t['calendar_days'],t['patient_days_raw'],t['patient_days'],'各床別先加總未取整值後取整',t['daily_status']]])
 sheet(wb,'人日耗用趨勢',['年月','耗用金額','平均床數','日曆天數','未取整人日','住院人日','每人日耗用','標準30日每床耗用','耗用範圍','狀態'],[[x['period'],x['used'],x.get('value'),x['calendar_days'],x['patient_days_raw'],x['patient_days'],x['per_patient_day'],x['standard30_per_bed'],x['care_scope_name'],x['daily_status']] for x in d['trend']])
 sheet(wb,'分類人日耗用',['分類','耗用金額','耗用占比%','每人日耗用','標準30日每床耗用'],[[x['category'],x['used'],x['usage_share'],x['per_patient_day'],x['standard30_per_bed']] for x in d['categories']])
 sheet(wb,'存貨期間彙總',['月數','期間','耗用合計','按天數加權平均存貨','期間週轉率','期間週轉天數','期間天數','口徑或不足原因'],[[n,','.join(x['periods']),x['used'],x['average'],x['turnover'],x['days'],x['calendar_days'],x['reason']] for n,x in d['window_totals'].items()])
 sheet(wb,'庫存風險矩陣',['品號','品名','期末存貨','週轉天數','庫存覆蓋天數','歷史推估低庫存覆蓋天數','風險狀態'],[[x[k] for k in ('code','name','ending','days','coverage','low_stock_days','state')] for x in d['risk_matrix']])
 sheet(wb,'自動計算檢查',['醫院','年月','可計算','可績效比較','具體原因'],[[x['hospital'],x['period'],x['calculable'],x['comparison_eligible'],x['reason']] for x in d['scope_availability']])
 sheet(wb,'分床別耗用',['床別','耗用金額','住院人日','每人日耗用','狀態'],[[x['department'],x['used'],x['patient_days'],x['per_patient_day'],x['daily_status']] for x in d['bed_breakdown']])
 tracking_sheet(wb,[d],sheet)
 from openpyxl.chart import LineChart,Reference,ScatterChart,Series
 ws=wb['人日耗用趨勢']
 if ws.max_row>1:
  chart=LineChart();chart.title='每住院人日耗用趨勢';chart.display_blanks='gap';chart.add_data(Reference(ws,min_col=7,min_row=1,max_row=ws.max_row),titles_from_data=True);chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=ws.max_row));ws.add_chart(chart,'L2')
 rs=wb['庫存風險矩陣']
 if rs.max_row>1:
  chart=ScatterChart();chart.title='庫存金額與週轉天數';chart.x_axis.title='週轉天數';chart.y_axis.title='庫存金額';series=Series(Reference(rs,min_col=3,min_row=2,max_row=rs.max_row),Reference(rs,min_col=4,min_row=2,max_row=rs.max_row),title='品項');series.graphicalProperties.line.noFill=True;series.marker.symbol='circle';series.marker.size=4;chart.series.append(series);rs.add_chart(chart,'I2')


def tracking_sheet(wb,reports,sheet):
 if '異常改善追蹤報表' in wb:del wb['異常改善追蹤報表']
 rows=[]
 for d in reports:
  reviews={(r['period'],r['code'],r['rule']):r for r in d.get('case_history',d['reviews'])}
  findings={(d['period'],f['code'],f['rule']):f for f in d['findings']}
  for key in sorted(set(reviews)|set(findings)):
   f=findings.get(key,{});r=reviews.get(key,{});details=r.get('details') or {};item=next((x for x in d['items'] if x['code']==key[1]),{})
   rows.append([d['hospital'],key[0],item.get('category',details.get('category',key[1])),key[1],f.get('name',item.get('name')),key[2],f.get('basis',details.get('basis')),__import__('app.inpatient',fromlist=['decimal']).decimal(f.get('value',details.get('difference'))),__import__('app.inpatient',fromlist=['decimal']).decimal(f.get('impact_amount',details.get('impact_amount'))),r.get('confirmed',False),r.get('cause','待確認'),details.get('responsible_unit'),r.get('owner'),r.get('measure'),r.get('due_date'),details.get('outcome','unassessed'),details.get('result'),details.get('completed_date'),__import__('decimal').Decimal(details['result_amount']) if details.get('result_amount') else None,r.get('status','待確認'),r.get('approved',False),r.get('revision'),r.get('updated_at'),d.get('selected_department') or '未標示'])
 return sheet(wb,'異常改善追蹤報表',['醫院','年月','分類','品號或分類','品名','規則','比較基準','異常數值','影響金額非損益','原因確認','原因','負責單位','負責人','改善措施','預計完成日期','結果類別','改善後結果','實際完成日期','人工核定成果參考金額','處理狀態','主管覆核','修訂版','更新時間','科別'],rows)


def overview_workbook(wb,reports,overview,sheet,kind):
 from datetime import datetime,timezone
 total=overview['total'];p=overview['period']
 sheet(wb,'管理總覽月報',['項目','值'],[['報表類型',{'group':'醫院合併管理總覽月報','single':'單院耗材管理月報','tracking':'異常改善追蹤報表'}[kind]],['年月',p],['科別',reports[0].get('selected_department') or '依各院月報'],['醫院',','.join(d['hospital'] for d in reports)],['月耗用總金額',total['used']],['每床月耗用',total['monthly_bed_consumption']],['當月登錄平均床數合計',total['total_beds']],['每床月耗用基準',total['baseline_monthly_bed_consumption']],['每床月耗用口徑',total['bed_month_status']],['當月住院人日',total['patient_days']],['未取整人日合計',total['patient_days_raw']],['全月報每住院人日耗用（參考）' if total['care_scope']=='hospital_reference' else '每住院人日耗用',total['per_patient_day']],['耗用異常影響金額',total['consumption_anomaly_impact']],['期末庫存',total['ending']],['整體週轉天數',total['days']],['實際改善達成率',total['improvement_rate']],['改善案件數',total['improvement_cases']],['實際改善件數',total['improved_cases']],['合理原因結案件數',total['reasonable_closed_cases']],['未改善件數',total['unresolved_count']],['計算口徑',total['note']],['狀態',total['reason']],['匯出時間',datetime.now(timezone.utc).isoformat()]])
 ws=wb['管理總覽月報']
 for label,key in [('本月底應完成件數','due_cases'),('同批案件改善完成件數','due_improved_cases'),('同批案件完成率%','due_completion_rate'),('逾期件數','overdue_cases'),('合理原因結案件數另列','reasonable_closed_count'),('應完成案件統計口徑','due_definition'),('比較基準金額','baseline_used'),('同範圍本月金額','comparable_current_used')]:
  v=total.get(key);ws.append([label,float(v) if hasattr(v,'as_tuple') else v])
 ws.append(['本月涵蓋醫院',','.join(overview['covered_hospitals'])]);ws.append(['增減比較同範圍醫院',','.join(total['comparison_hospitals'])]);ws.append(['增減差額',float(total['change']['amount']) if total['change']['amount'] is not None else None])
 sheet(wb,'資料完整狀態',['醫院','月份','用途','月報','床數','具體原因'],[[x['hospital'],x['period'],x['role'],x['report'],x['beds'],x['reason']] for x in overview['data_status']])
 sheet(wb,'分類差異',['分類','本月金額','占比%','基準金額','增減金額','增減%','每人日耗用'],[[x['category'],x['used'],x['usage_share'],x['baseline_used'],x['change']['amount'],x['change']['pct'],x['per_patient_day']] for x in overview['categories']])
 sheet(wb,'品項與差異原因',['醫院','品號','品名','單位','耗用金額','耗用數量','每人日耗用量','基準金額','增減金額','數量影響金額','單位成本影響金額','拆解合計金額','不可比原因'],[[d['hospital'],x['code'],x['name'],x['unit'],x['used'],x['qty'],x['per_patient_day_qty'],x['baseline_used'],x['change']['amount'],x['selected_factor'].get('quantity'),x['selected_factor'].get('unit_cost'),x['selected_factor'].get('total'),x['selected_factor'].get('reason')] for d in reports for x in d['items']])
 sheet(wb,'案件成果口徑',['醫院','應完成件數','改善完成件數','完成率%','逾期件數','合理原因結案件數','口徑'],[[d['hospital'],d['total']['due_cases'],d['total']['due_improved_cases'],d['total']['due_completion_rate'],d['total']['overdue_cases'],d['total']['reasonable_closed_count'],d['total']['due_definition']] for d in reports])
 if kind=='single':
  single=reports[0];ws=wb['管理總覽月報'];ws.title='單院耗材管理月報'
  for label,key in [('標準30日每床耗用','standard30_per_bed'),('當月每床月耗用自動對應範圍','monthly_per_bed')]:
   value=single['total'][key];ws.append([label,float(value) if value is not None else None])
  sheet(wb,'前三月加權基準',['項目','值'],[['基準月份',','.join(single['daily_baseline3']['periods'])],['前三月耗用合計',single['daily_baseline3']['used']],['前三月住院人日',single['daily_baseline3']['patient_days']],['前三月每人日耗用基準',single['daily_baseline3']['per_patient_day']],['前三月偏差率%',single['total']['daily_change3']['pct']],['狀態',single['total']['daily_status']]])
 sheet(wb,'各院月報',['醫院','耗用金額','耗用增減%','高庫存／本月無耗用金額','期末庫存','平均庫存','週轉率','週轉天數','住院人日','未取整人日','每人日耗用','標準30日每床耗用','照護範圍','狀態','啟用版本','門檻版本'],[[d['hospital'],d['total']['used'],d['total']['change']['pct'],d['total']['high_low_inventory_amount'],d['total']['ending'],d['total']['average'],d['total']['turnover'],d['total']['days'],d['total']['patient_days'],d['total']['patient_days_raw'],d['total']['per_patient_day'],d['total']['standard30_per_bed'],d['total']['care_scope_name'],d['total']['daily_status'],d['versions'],d['rule_version']] for d in reports])
 sheet(wb,'各院分類',['醫院','分類','耗用金額','耗用占比%','每人日耗用','期末存貨','週轉天數'],[[d['hospital'],x['category'],x['used'],x['usage_share'],x['per_patient_day'],x['ending'],x['days']] for d in reports for x in d['categories']])
 sheet(wb,'管理合計月趨勢',['月份','耗用','期末存貨','平均存貨','週轉天數','住院人日','每人日耗用','登錄平均床數合計','每床月耗用'],[[x.get(k) for k in ('period','used','ending','average','days','patient_days','per_patient_day','total_beds','monthly_bed_consumption')] for x in overview['trend']])
 sheet(wb,'管理摘要',['醫院','摘要'],[['合計',x] for x in overview['summary']]+[[d['hospital'],x] for d in reports for x in d['summary']])
 tracking_sheet(wb,reports,sheet)
 sheet(wb,'床別與人日',['醫院','年月','床別','原始平均床數','當月天數','未取整人日','住院人日','來源'],[[d['hospital'],d['period'],x['department'],x.get('value'),x['calendar_days'],x['patient_days_raw'],x['patient_days'],x['source']] for d in reports for x in d['census']['records']])
 from .dashboard_reports import extend
 extend(wb,reports,overview,sheet)
 from openpyxl.chart import LineChart,Reference
 ws=wb['管理合計月趨勢']
 if ws.max_row>1:
  for col,title,anchor in [(2,'合計耗用趨勢','I2'),(7,'合計每人日耗用','I18'),(3,'合計存貨','I34')]:
   chart=LineChart();chart.title=title;chart.display_blanks='gap';chart.add_data(Reference(ws,min_col=col,min_row=1,max_row=ws.max_row),titles_from_data=True);chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=ws.max_row));ws.add_chart(chart,anchor)
