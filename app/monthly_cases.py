"""月異常確認、原案連結及五種進度。所有寫入保留逐品項歷程。"""
import calendar
from copy import deepcopy
from datetime import datetime, date
from decimal import Decimal as D
from io import BytesIO
from zoneinfo import ZoneInfo
from fastapi import HTTPException, Depends, Response
from psycopg.types.json import Jsonb
from openpyxl import Workbook
from . import auth, db, tracking, engine
from .analysis_grid import threshold
from .inventory_rules import flags
from . import monthly_workflow as wf

TITLES = {'usage': '耗用異常', 'purchase': '採購單價異常', 'high': '高庫存低耗用', 'low': '低庫存風險', 'idle': '無耗用仍有庫存'}
DONE = {'合理原因已核准', '已驗證改善'}  # 僅供舊資料辨識，不是新流程進度。
CAUSES = list(dict.fromkeys([*wf.LEGACY_CAUSES, *(v for values in wf.CAUSE_OPTIONS.values() for v in values)]))
BASES = {'mean3': '近三個完整月平均', 'previous': '上月', 'year': '去年同月'}
REVIEWERS = ('accountant', 'manager', 'admin')
FILLERS = ('frontline', 'admin')
TRACKING_START = '2026-09'
SYNC_VERSION = '0111-ui-2'

def sync_token(c):
    from .performance import source_revision
    return SYNC_VERSION + ':' + source_revision(c)

def sync_fresh(c, hospital, period, base, token, existing_only=False):
    row = c.execute('SELECT source_token, existing_only FROM monthly_sync_state WHERE hospital=%s AND period=%s AND basis=%s', (hospital, period, base)).fetchone()
    return bool(row and row['source_token'] == token and (existing_only or not row['existing_only']))

def source_index(d):
    if '_performance_case_sources' not in d:
        index = {}
        for v in d.get('source_rows', []):
            if v.get('valid'):
                index.setdefault((v['code'], v['name'], v['unit']), []).append(v)
        d['_performance_case_sources'] = index
    return d['_performance_case_sources']


DOMAINS = {'usage': ('usage',), 'inventory': ('high', 'low', 'idle'), 'purchase': ('purchase',)}


def in_domain(kind, domain):
    if domain and domain not in DOMAINS:
        raise HTTPException(422, '分析項目不正確')
    return not domain or kind in DOMAINS[domain]


def in_view(t, x, view, period):
    state = wf.status(x)
    if view == 'completed':
        return state == '已完成'
    if view in ('confirm','follow') and x.get('tracking_excluded'):
        return False
    if view == 'follow':
        return bool(wf.improvement(x) and not x.get('linked_to') and state != '已完成')
    if view == 'confirm':
        return bool(state != '已完成' and not x.get('linked_to') and not wf.improvement(x) and x.get('comparison_active', x.get('active', False)))
    return True


def matches_state(x, state, view):
    if not state:
        return True
    current = wf.status(x)
    if view == 'follow' and state == '待審核':
        return current in ('待審核', '待確認結案')
    if view == 'follow' and state == '改善中':
        return current in ('改善中', '待填原因')
    return current == state


def own_todo(x, user):
    state = wf.status(x)
    return bool(user['role'] in FILLERS and (state == '待填原因' or state in ('改善中','待審核') and wf.improvement(x) and not x.get('needs_reconfirm') and x.get('return_stage') != 'reason') or user['role'] in REVIEWERS and state in ('待審核','待確認結案'))


def export_date(value):
    if not value:
        return None
    try:
        d = datetime.fromisoformat(str(value).replace('Z', '+00:00'))
        return (d.replace(tzinfo=wf.TZ) if d.tzinfo is None else d.astimezone(wf.TZ)).replace(tzinfo=None)
    except (ValueError, TypeError):
        return None


def scoped(u, h, write=False):
    auth.authorize(u, h, write=write)


def key(x):
    return x['code'] + '|' + x.get('unit', '')


def basis(d):
    if len(d['baseline']) == 3:
        return 'mean3'
    return 'previous' if d['baseline'] == [engine.previous(d['period'])] else 'year'


def eligible(x, d, kind):
    if kind == 'idle':
        f = flags(x, d['rules'])
        return bool(f['idle'] and f['persistent'])
    return threshold(x, d['rules']) if kind == 'usage' else bool(x.get('purchase_alert')) if kind == 'purchase' else flags(x, d['rules'])[kind]


def snapshot(x, d, kind):
    from .analytics import rule_for
    r = rule_for(d['rules'], x['category']); f = flags(x, d['rules']); ch = x['change']
    current = d['total']; trend = {v['period']: v for v in d.get('trend', [])}
    prior = [trend.get(p, {}) for p in d['baseline']]
    same_scope = bool(current.get('care_scope')) and all(v.get('care_scope') == current['care_scope'] and v.get('per_patient_day') is not None for v in prior)
    den = sum((D(str(v['patient_days'])) for v in prior), D(0)) if same_scope else None
    item_history = source_index(d).get((x['code'], x['name'], x['unit']), [])
    histories = [v for v in item_history if v['period'] in d['baseline']]
    evidence_periods = x.get('purchase_periods',[]) if kind=='purchase' else d['baseline']
    if kind=='low':
        evidence_periods=[engine.previous(d['period'])]
        for _ in range(2):evidence_periods.append(engine.previous(evidence_periods[-1]))
    evidence_rows = [v for v in item_history if v['period'] in evidence_periods]
    if kind=='high':evidence_rows=[]
    if kind=='idle':evidence_rows=[v for v in item_history if v['period']<d['period']][-12:]
    normalized = D(str(x['per_patient_day'])) * 30 if x.get('per_patient_day') is not None else None
    normalized_base = sum((engine.calc(v)['used'] for v in histories), D(0)) / den * 30 if den and len(histories) == len(d['baseline']) else None
    previous_rows = [v for v in item_history if v['period'] == engine.previous(d['period'])]
    previous_used = engine.calc(previous_rows[0])['used'] if len(previous_rows) == 1 else None
    names = ('change_pct', 'change_amount') if kind == 'usage' else ('purchase_pct', 'purchase_amount') if kind == 'purchase' else ('high_days', 'severe_days', 'low_days', 'idle_warn_days', 'idle_days')
    return tracking.plain(dict(code=x['code'], name=x['name'], unit=x['unit'], category=x['category'], used=x['used'], baseline=x.get('baseline_used'), difference=ch['amount'], pct=ch['pct'], previous_change=engine.delta(x['used'], previous_used), quantity=x['qty'], baseline_quantity=(sum((engine.calc(v)['qty'] for v in histories),D(0))/len(d['baseline']) if len(histories)==len(d['baseline']) and d['baseline'] else None), ending_quantity=next((D(str(v['end_q'])) for v in item_history if v['period']==d['period']),None), ending=x['ending'], days=x['days'], available=x.get('low_stock_days'), severity=f['severity'], long_term=f['long_term'], persistent=f['persistent'], idle_days=x.get('idle_observed_days'), valid=(x['used']>=0 and x['qty']>=0) if kind=='usage' else f['valid'], basis=basis(d), baseline_periods=d['baseline'], threshold={k:r[k] for k in names}, sources=[{k:v.get(k) for k in ('batch_id', 'version', 'row', 'period', 'filename', 'file_id')} for v in [x,*evidence_rows]], rule_version=d['rule_version'], standard30=normalized, standard30_baseline=normalized_base, standard30_change=engine.delta(normalized, normalized_base), patient_days=current.get('patient_days'), care_scope=current.get('care_scope'), scope_note=current.get('daily_status'), purchase=x.get('purchase'), purchase_basis=x.get('purchase_basis'), purchase_pct=x.get('purchase_delta', {}).get('pct'), purchase_impact=x.get('purchase_impact'), purchase_periods=x.get('purchase_periods', [])))


def event(c, t, u, action, detail):
    tracking.event(c, t['hospital'], '月案件／' + action, dict(monthly_case_id=t['id'], **detail), u['id'] if u else None)


def sync(c, d, existing_only=False):
    """新觀測捕捉當時基準。切換比較不覆寫既有案件原始依據。"""
    if d.get('hospital') in (None, 'ALL') or d['period'] < TRACKING_START:
        return
    h, p = d['hospital'], d['period']; db.write_lock(c)
    c.execute('SELECT pg_advisory_xact_lock(hashtextextended(%s,0))', ('monthly:' + h + ':' + p,))
    token = SYNC_VERSION + ':' + d['_performance_source'] if d.get('_performance_source') else sync_token(c)
    base = basis(d)
    excluded_identities=exclusion_index(c)
    if sync_fresh(c, h, p, base, token, existing_only):
        return
    for kind in TITLES:
        old = c.execute('SELECT * FROM monthly_cases WHERE hospital=%s AND period=%s AND kind=%s FOR UPDATE', (h, p, kind)).fetchone()
        previous = (old or {}).get('body', {}).get('items', {}); items = deepcopy(previous)
        if existing_only and not old:
            continue
        for row in d['items']:
            k = key(row); active = eligible(row, d, kind)
            if k not in items and (h,row['code'],row['name'],row['unit']) in excluded_identities:continue
            if (not active or existing_only) and k not in items:
                continue
            data = snapshot(row, d, kind)
            fields = {'usage': ('used','baseline','difference','pct','quantity','standard30','standard30_baseline','basis','baseline_periods'), 'purchase': ('purchase','purchase_basis','purchase_pct','purchase_impact','purchase_periods'), 'high': ('ending','days','used','quantity','severity'), 'low': ('ending','available','quantity'), 'idle': ('ending','used','quantity','long_term','idle_days')}[kind]
            facts = {k:data[k] for k in ('code','name','unit','category','valid','threshold','rule_version',*fields)}
            facts['sources'] = data['sources'] if kind == 'usage' else data['sources'][:1]
            fp = tracking.digest(facts)
            if k not in items:
                x = dict(status='待填原因', workflow_status='待填原因', active=active, present=True, data=data, source_data=data, fingerprint=fp, created_at=datetime.now(wf.TZ).isoformat())
                wf.record(x, 'created', {})
            else:
                x = wf.normalize(items[k]); x['present'] = True
                if not items[k].get('workflow_status') and not x.get('history'):
                    wf.record(x, 'migrated', {})
                # 耗用案件用建立時的基準覆核來源，不因下拉選單改變狀態。
                saved_basis = x['data'].get('basis') or ('mean3' if len(x['data'].get('baseline_periods', [])) == 3 else 'previous')
                if kind != 'usage' or saved_basis == basis(d):
                    if 'source_data' not in x:
                        x.update(source_data=data, fingerprint=fp)
                    if x.get('fingerprint') != fp:
                        x.update(source_changed=True, source_data=data, fingerprint=fp)
                        if wf.status(x) != '已完成':
                            x.update(workflow_status='待填原因', needs_reconfirm=True, return_stage='reason')
                        wf.record(x, 'source_changed', {})
                    x['active'] = active
            items[k] = x
        codes = {key(x) for x in d['items']}
        for k, x in items.items():
            if k not in codes and x.get('present', True):
                x.update(present=False, active=False, source_changed=True)
                if wf.status(x) != '已完成':
                    x.update(workflow_status='待填原因', needs_reconfirm=True, return_stage='reason')
                wf.record(x, 'source_missing', {})
        if not old and not items:
            continue
        body = {**((old or {}).get('body') or {}), 'items': items, 'workflow_version': 99}
        if old:
            if old['body'] != body:
                c.execute('UPDATE monthly_cases SET body=%s,revision=revision+1,updated_at=now() WHERE id=%s', (Jsonb(body), old['id']))
        else:
            t = c.execute('INSERT INTO monthly_cases(hospital,period,kind,body) VALUES(%s,%s,%s,%s) RETURNING *', (h, p, kind, Jsonb(body))).fetchone()
            event(c, t, None, '自動建立', dict(items=list(items)))

    c.execute('INSERT INTO monthly_sync_state(hospital,period,basis,source_token,existing_only) VALUES(%s,%s,%s,%s,%s) ON CONFLICT(hospital,period,basis) DO UPDATE SET source_token=excluded.source_token,existing_only=excluded.existing_only', (h,p,base,token,existing_only))


def display_status(x):
    return wf.status(x)


def compact(t):
    items = {k:{**wf.normalize(v), 'display_status': wf.status(v), 'next_owner': wf.next_owner(v), 'overdue_days': wf.overdue(v)} for k,v in t['body']['items'].items()}
    return tracking.plain(dict(id=t['id'], hospital=t['hospital'], period=t['period'], kind=t['kind'], title=t['period'].replace('-', '') + '_' + TITLES[t['kind']], revision=t['revision'], items=items, counts={s:sum(x['display_status'] == s for x in items.values()) for s in wf.STATES}, complete=bool(items) and all(wf.status(x) == '已完成' for x in items.values()), updated_at=t['updated_at']))


def require_plan(b):
    result = {k:str(b.get(k, '')).strip() for k in ('owner', 'deadline', 'plan', 'target')}
    if not result['deadline'] or not result['plan']:
        raise HTTPException(422, '請填預計完成日及改善措施')
    try:
        date.fromisoformat(result['deadline'])
    except ValueError:
        raise HTTPException(422, '請填有效的預計完成日')
    return result


def act(c, t, b, u):
    scoped(u, t['hospital'], write=True)
    if b.get('revision') != t['revision']:
        raise HTTPException(409, '案件已更新，請重新開啟')
    keys = list(dict.fromkeys(b.get('keys', []))); original = t['body']['items']; action = b.get('action'); note = str(b.get('note', '')).strip()
    if not keys or any(k not in original for k in keys):
        raise HTTPException(422, '請選擇此案品項')
    if not note:
        raise HTTPException(422, '請填說明')
    if action not in ('reason', 'approve', 'approve_follow', 'reject', 'progress', 'submit', 'finish', 'reopen', 'link', 'unlink', 'group'):
        raise HTTPException(422, '操作不正確')
    if action in ('approve', 'approve_follow', 'reject', 'finish', 'reopen', 'unlink') and u['role'] not in REVIEWERS:
        raise HTTPException(403, '需要會計、主管或系統管理員審核')
    if action in ('reason', 'progress', 'submit', 'link') and u['role'] not in FILLERS:
        raise HTTPException(403, '需要前線人員或系統管理員填報')
    if action == 'group' and u['role'] not in FILLERS + REVIEWERS:
        raise HTTPException(403, '沒有加入改善案的權限')
    if len({wf.status(original[k]) for k in keys}) != 1:
        raise HTTPException(422, '請分開選擇相同進度的品項')
    if action == 'approve' and len({original[k].get('submission') for k in keys}) != 1:
        raise HTTPException(422, '原因與改善結果請分開審核')
    # 先驗證完整批次，以免任何一項失敗造成部分變更。
    items = deepcopy(original); now = datetime.now(wf.TZ).isoformat()
    for k in keys:
        x = wf.normalize(items[k]); s = wf.status(x)
        if x.get('linked_to') and action != 'unlink':
            raise HTTPException(422, '此提醒已連結原改善案，請至原案辦理')
        allowed = {'reason': ('待填原因',), 'approve': ('待審核',), 'approve_follow': ('待審核',), 'reject': ('待審核', '待確認結案'), 'progress': ('改善中', '待審核'), 'submit': ('改善中', '待審核'), 'finish': ('待確認結案',), 'reopen': ('已完成',), 'link': ('待填原因',), 'unlink': wf.STATES, 'group': tuple(s for s in wf.STATES if s != '已完成')}
        if s not in allowed[action]:
            raise HTTPException(422, '目前進度不允許此操作，請重新讀取案件')
        if action in ('reason', 'approve', 'approve_follow') and (not x.get('present') or not x.get('source_data', x['data']).get('valid')):
            raise HTTPException(422, '資料不足或無效，請先修正來源')
        if action == 'reason':
            if x.get('return_stage') == 'result':
                raise HTTPException(422, '請補充改善結果後再次送審')
            cause = b.get('cause_category')
            if cause not in wf.CAUSE_OPTIONS[t['kind']] and cause != x.get('cause_category'):
                raise HTTPException(422, '請選擇此異常類型適用的原因類別')
            choice = b.get('follow_up')
            if choice not in (True, False, 'yes', 'no', 'true', 'false'):
                raise HTTPException(422, '請明確選擇是否需要改善')
            follow = choice in (True, 'yes', 'true')
            plan = require_plan({**b, 'owner':str(b.get('owner') or x.get('owner') or u['name']), 'target':b.get('target') or x.get('target') or ''}) if follow else dict(owner=None, deadline=None, plan=None, target=None)
            x.update(**plan, cause_approved=False, approved_cause=None, follow_up=follow, follow_up_completed=False, cause_category=cause, note=note, evidence=str(b.get('evidence', x.get('evidence') or '')).strip(), submitted_at=now, submission='reasonable', submitted_snapshot=deepcopy(x.get('source_data', x['data'])), needs_reconfirm=False, return_stage=None, workflow_status='待審核', operator=u['name'], operation_date=now)
        elif action in ('approve', 'approve_follow'):
            if x.get('needs_reconfirm'):
                raise HTTPException(409, '來源已變更，需前線重新確認')
            if action == 'approve_follow' and not x.get('follow_up'):
                raise HTTPException(422, '請退回前線補填改善計畫，再確認追蹤')
            if action == 'approve' and x.get('follow_up'):
                raise HTTPException(422, '已填改善計畫，請使用確認並追蹤改善')
            x.update(cause_approved=True, approved_cause=dict(category=x['cause_category'], note=x['note'], evidence=x.get('evidence'), at=now), reviewer=u['name'], review_date=now, review_note=note, workflow_status='改善中' if x.get('follow_up') else '已完成')
            if not x.get('follow_up'):
                x.update(completed_at=now, completion_kind='免追蹤')
        elif action == 'reject':
            x.update(return_stage='result' if s == '待確認結案' else 'reason', workflow_status='改善中' if s == '待確認結案' else '待填原因', review_note=note, reviewer=u['name'], review_date=now)
        elif action == 'progress':
            if not wf.improvement(x) or x.get('needs_reconfirm') or x.get('return_stage') == 'reason':
                raise HTTPException(422, '請先補正原因與來源，再更新改善進度')
            x.update(progress_note=note, progress_at=now, operator=u['name'], operation_date=now, progress_evidence=str(b.get('evidence', '')).strip())
        elif action == 'submit':
            if not wf.improvement(x) or x.get('needs_reconfirm') or x.get('return_stage') not in (None, 'result'):
                raise HTTPException(422, '請先送出改善計畫並補正原因與來源')
            evidence = str(b.get('evidence', '')).strip()
            x.update(workflow_status='待確認結案', result_note=note, result_evidence=evidence, result_at=now, result_period=b.get('result_period'), result_data=b.get('_result_by_key', {}).get(k), return_stage=None, operator=u['name'], operation_date=now)
        elif action == 'finish':
            if x.get('needs_reconfirm') or not x.get('present') or not x.get('source_data', x['data']).get('valid'):
                raise HTTPException(422, '來源已變更或無效，請先退回補正')
            if not x.get('result_note'):
                raise HTTPException(422, '缺少改善結果，不能結案')
            if not x.get('cause_approved'):
                x.update(cause_approved=True, approved_cause=dict(category=x['cause_category'], note=x['note'], evidence=x.get('evidence'), at=now))
            x.update(workflow_status='已完成', follow_up_completed=True, completed_at=now, completion_kind='改善完成', review_note=note, reviewer=u['name'], review_date=now, verified_data=x.get('result_data'))
        elif action == 'reopen':
            x.pop('linked_to', None)
            x.update(workflow_status='改善中' if x.get('cause_approved') and x.get('follow_up') else '待填原因', follow_up_completed=False, completed_at=None, review_note=note, reviewer=u['name'], review_date=now)
        elif action == 'unlink':
            if not x.get('linked_to'):
                raise HTTPException(422, '此提醒沒有連結原案')
            x.pop('linked_to'); x.update(workflow_status='待填原因', note=note, operator=u['name'], operation_date=now)
        elif action == 'group':
            target_id, target_key = b.get('target_case_id'), b.get('target_key')
            target = c.execute('SELECT * FROM monthly_cases WHERE id=%s FOR UPDATE', (target_id,)).fetchone()
            root = (target or {}).get('body', {}).get('items', {}).get(target_key)
            domain = lambda kind: 'inventory' if kind in ('high', 'low', 'idle') else kind
            if (not target or not root or target['hospital'] != t['hospital'] or domain(target['kind']) != domain(t['kind'])
                    or target['period'] > t['period'] or (target_id == t['id'] and target_key == k)
                    or root.get('linked_to') or root.get('group_root') or not root.get('follow_up')
                    or not root.get('cause_approved') or wf.status(root) == '已完成'):
                raise HTTPException(422, '請選同院、同分析類別、同月或更早且未完成的主改善案；不可選自己或成員案')
            # 保留個別品項的審核、目標與進度，不以主案完成代替成員驗證。
            if any(v.get('group_root', {}).get('case_id') == t['id'] and v.get('group_root', {}).get('key') == k
                   for row in c.execute('SELECT body FROM monthly_cases WHERE hospital=%s', (t['hospital'],)).fetchall()
                   for v in row['body']['items'].values()):
                raise HTTPException(422, '已有成員的主改善案不可再加入其他主案')
            x.update(group_root=dict(case_id=target_id, key=target_key, period=target['period'], name=root['data']['name']), group_note=note)
        elif action == 'link':
            target_id = b.get('target_case_id'); target_key = b.get('target_key')
            target = c.execute('SELECT * FROM monthly_cases WHERE id=%s FOR UPDATE', (target_id,)).fetchone()
            root = (target or {}).get('body', {}).get('items', {}).get(target_key)
            if not target or not root or target['hospital'] != t['hospital'] or target['kind'] != t['kind'] or target['period'] >= t['period'] or target_key != k or root['data']['name'] != x['data']['name'] or root.get('linked_to') or not root.get('follow_up') or not root.get('cause_approved') or wf.status(root) == '已完成':
                raise HTTPException(422, '只能連結同院、同類型、同品號品名單位且未完成的前期改善案')
            x.update(linked_to=dict(case_id=target_id, key=target_key, period=target['period']), note=note)
        x['status'] = x['workflow_status']; wf.record(x, action, u, now); items[k] = x
    body = {**t['body'], 'items': items, 'workflow_version': 99}
    event(c, t, u, action, dict(keys=keys, note=note, before=original, after=items))
    c.execute('UPDATE monthly_cases SET body=%s,revision=revision+1,updated_at=now() WHERE id=%s', (Jsonb(body), t['id']))
    db.audit(c, u['id'], 'monthly_' + action, t['hospital'], t['id'], dict(keys=keys))
    return dict(id=t['id'])


def decorate(rows, reports):
    result = [compact(t) for t in rows]; index = {(t['id'], k):x for t in result for k,x in t['items'].items()}
    current_rows = {h:{key(x):x for x in d['items']} for h,d in reports.items()}
    latest = {}; candidates = {}
    for old in result:
        for ok, ox in old['items'].items():
            if ox.get('cause_approved') and ox.get('follow_up') and not ox.get('linked_to') and wf.status(ox) != '已完成':
                candidates.setdefault((old['hospital'],old['kind'],ok,ox['data']['name']), []).append(dict(case_id=old['id'],key=ok,period=old['period']))
    for t in result:
        d = reports.get(t['hospital']); current = current_rows.get(t['hospital'], {})
        for k, x in t['items'].items():
            row = current.get(k)
            if row and row['name'] == x['data']['name']:
                cache_key = (t['hospital'], t['kind'], k, row['name'])
                if cache_key not in latest:
                    latest[cache_key] = snapshot(row, d, t['kind'])
                x['latest_data'] = latest[cache_key]; x['latest_period'] = d['period']
                x['comparison_active'] = eligible(row, d, t['kind'])
            else:
                x['latest_data'] = None; x['comparison_active'] = False
            x['link_candidates'] = [v for v in candidates.get((t['hospital'],t['kind'],k,x['data']['name']), []) if v['period'] < t['period']]
            if x.get('linked_to'):
                root = index.get((x['linked_to']['case_id'], x['linked_to']['key']))
                if root:
                    x.update(display_status=wf.status(root), next_owner=wf.next_owner(root), overdue_days=wf.overdue(root), linked_status=wf.status(root))
    return result


def exclusion_index(c):
    return {(v['hospital'],v['code'],v['name'],v['unit']):v['excluded_at'].isoformat() for v in c.execute('SELECT * FROM tracking_exclusions WHERE active').fetchall()}


def register(app, report, jsonable, sheet):
    @app.get('/api/monthly-cases/exclusions')
    def exclusions(hospitals:str,u=Depends(auth.current)):
        hs=list(dict.fromkeys(hospitals.split(',')))
        for h in hs:scoped(u,h)
        with db.connect() as c:
            return jsonable(c.execute('SELECT id,hospital,code,name,unit,reason,excluded_at,operator FROM tracking_exclusions WHERE hospital=ANY(%s) AND active ORDER BY hospital,code',(hs,)).fetchall())

    @app.post('/api/monthly-cases/exclusions')
    def exclude(body:dict,u=Depends(auth.current)):
        if u['role'] not in REVIEWERS:raise HTTPException(403,'免列入追蹤須由會計、主管或系統管理員設定')
        note=str(body.get('reason','')).strip()
        if not note:raise HTTPException(422,'請填免列入追蹤或恢復原因')
        with db.connect() as c:
            db.write_lock(c)
            if body.get('restore_id'):
                row=c.execute('SELECT * FROM tracking_exclusions WHERE id=%s FOR UPDATE',(body['restore_id'],)).fetchone()
                if not row:raise HTTPException(404,'設定不存在')
                scoped(u,row['hospital'],write=True)
                c.execute('UPDATE tracking_exclusions SET active=false WHERE id=%s',(row['id'],))
                c.execute('DELETE FROM monthly_sync_state WHERE hospital=%s',(row['hospital'],))
                db.audit(c,u['id'],'tracking_exclusion_restore',row['hospital'],row['id'],dict(reason=note))
                return dict(restored=True)
            selected=body.get('items')
            if not isinstance(selected,list) or not 1<=len(selected)<=200:raise HTTPException(422,'請選1至200個品項')
            identities=set()
            for item in selected:
                t=c.execute('SELECT * FROM monthly_cases WHERE id=%s FOR UPDATE',(item.get('case_id'),)).fetchone()
                if not t:raise HTTPException(404,'案件不存在')
                scoped(u,t['hospital'],write=True)
                if t['revision']!=item.get('revision'):raise HTTPException(409,'案件已更新，請重新整理')
                x=t['body']['items'].get(item.get('key'))
                if not x:raise HTTPException(422,'品項不存在')
                d=x['data'];identities.add((t['hospital'],d['code'],d['name'],d['unit']))
            for h,code,name,unit in identities:
                rid=c.execute('INSERT INTO tracking_exclusions(hospital,code,name,unit,reason,operator) VALUES(%s,%s,%s,%s,%s,%s) ON CONFLICT(hospital,code,name,unit) DO UPDATE SET reason=excluded.reason,operator=excluded.operator,excluded_at=now(),active=true RETURNING id',(h,code,name,unit,note,u['id'])).fetchone()['id']
                db.audit(c,u['id'],'tracking_exclusion_set',h,rid,dict(code=code,name=name,unit=unit,reason=note))
        return dict(excluded=len(identities))

    @app.get('/api/monthly-cases')
    def listing(hospitals:str, month:str, records:bool=False, baseline:str='mean3', u=Depends(auth.current)):
        hs = list(dict.fromkeys(hospitals.split(',')))
        if not hs or len(hs) > 50 or baseline not in BASES:
            raise HTTPException(422, '查詢範圍或比較基準不正確')
        p = __import__('app.excel', fromlist=['period']).period(month)
        for h in hs:
            scoped(u, h)
        with db.connect() as c:
            reports = {h:report(c, h, p, baseline=baseline, department='__all__') for h in hs} if p >= TRACKING_START else {}
            rows = c.execute('SELECT * FROM monthly_cases WHERE hospital=ANY(%s) AND period>=%s AND period<=%s ORDER BY period,id', (hs, TRACKING_START, p)).fetchall()
            legacy_scopes = {(t['hospital'],t['period']) for t in rows if any(not v.get('workflow_status') or not v.get('history') for v in t['body']['items'].values())}
            for h, period in legacy_scopes:
                c.execute('DELETE FROM monthly_sync_state WHERE hospital=%s AND period=%s', (h,period))
            token = sync_token(c)
            needed = {(t['hospital'], t['period'], v['data'].get('basis', 'previous')) for t in rows for v in t['body']['items'].values()}
            # Prepare reports before any writes, so the report cache remains usable.
            pending = []
            for h, period, base in sorted(needed):
                base = base if base in BASES else 'previous'
                if not sync_fresh(c, h, period, base, token, existing_only=True):
                    d = reports.get(h) if period == p and base == baseline else report(c,h,period,baseline=base,department='__all__')
                    pending.append(d)
            for d in reports.values():
                sync(c, d)
            for d in pending:
                sync(c, d, existing_only=True)
            rows = c.execute('SELECT * FROM monthly_cases WHERE hospital=ANY(%s) AND period>=%s AND period<=%s ORDER BY period,id', (hs, TRACKING_START, p)).fetchall()
            cases = decorate(rows, reports)
            excluded=exclusion_index(c)
            for t in cases:
                for k,x in t['items'].items():x['tracking_excluded']=excluded.get((t['hospital'],x['data']['code'],x['data']['name'],x['data']['unit']))
        return jsonable(dict(cases=cases, tracking_start=TRACKING_START, period=p, baseline=baseline, baseline_label=BASES[baseline], cause_options=wf.CAUSE_OPTIONS, states=wf.STATES, case_count=sum(1 for t in cases for x in t['items'].values() if in_view(t,x,'confirm',p) or in_view(t,x,'follow',p)), data_gaps=[dict(hospital=h, missing=d.get('baseline_options', {}).get(baseline, {}).get('missing', []), uncomparable_items=sum(x['change']['pct'] is None for x in d['items']), current_missing=not bool(d['items'])) for h,d in reports.items()]))

    def execute_action(c, cid, body, u):
        db.write_lock(c); t = c.execute('SELECT * FROM monthly_cases WHERE id=%s', (cid,)).fetchone()
        if not t:
            raise HTTPException(404, '案件不存在')
        scoped(u, t['hospital'], write=True)
        db.scope_lock(c, t['hospital'], t['period']); db.check_open(c, t['hospital'], t['period'])
        for base in {v['data'].get('basis', 'previous') for v in t['body']['items'].values()}:
            sync(c, report(c, t['hospital'], t['period'], baseline=base if base in BASES else 'previous', department='__all__'), existing_only=True)
        t = c.execute('SELECT * FROM monthly_cases WHERE id=%s FOR UPDATE', (cid,)).fetchone()
        if body.get('action') == 'submit':
            rp = __import__('app.excel', fromlist=['period']).period(str(body.get('result_period', '')))
            if rp < t['period'] or rp > datetime.now(wf.TZ).strftime('%Y-%m'):
                raise HTTPException(422, '改善結果月份不可早於原始月份或晚於目前月份')
            d = report(c, t['hospital'], rp, baseline='mean3', department='__all__')
            # 多選時每個品項各自保存改善數據，不能共用第一項。
            values = {key(x): snapshot(x,d,t['kind']) for x in d['items']}
            body = {**body, '_result_by_key':values}
        return act(c, t, body, u)


    @app.post('/api/monthly-cases-batch')
    def batch_action(body:dict, u=Depends(auth.current)):
        operations = body.get('operations')
        if not isinstance(operations, list) or not 1 <= len(operations) <= 200:
            raise HTTPException(422, '請選擇 1 至 200 個品項')
        seen = set()
        with db.connect() as c:
            db.write_lock(c)
            # 全批先驗證原始版本，再逐筆套用；任一失敗整批回復。
            for op in operations:
                if not isinstance(op, dict) or not isinstance(op.get('case_id'), int) or len(op.get('keys', [])) != 1:
                    raise HTTPException(422, '批次品項格式不正確')
                identity = (op['case_id'], op['keys'][0])
                if identity in seen:
                    raise HTTPException(422, '批次品項不可重複')
                seen.add(identity)
                t = c.execute('SELECT * FROM monthly_cases WHERE id=%s FOR UPDATE', (op['case_id'],)).fetchone()
                if not t:
                    raise HTTPException(404, '案件不存在')
                scoped(u, t['hospital'], write=True)
                if t['revision'] != op.get('revision'):
                    raise HTTPException(409, '所選案件已更新，請重新整理後重選')
            for op in operations:
                t = c.execute('SELECT revision FROM monthly_cases WHERE id=%s', (op['case_id'],)).fetchone()
                execute_action(c, op['case_id'], {**op, 'revision': t['revision']}, u)
        return dict(updated=len(operations))

    @app.post('/api/monthly-cases/{cid}')
    def action(cid:int, body:dict, u=Depends(auth.current)):
        with db.connect() as c:
            return execute_action(c, cid, body, u)

    @app.get('/api/monthly-statistics')
    def statistics(hospitals:str, year:int, quarter:int=0, kind:str='', u=Depends(auth.current), domain:str='', month:str|None=None):
        in_domain('usage', domain)
        hs = list(dict.fromkeys(hospitals.split(',')))
        if not hs or len(hs) > 50 or not 1912 <= year <= 9998 or quarter not in range(5) or kind and kind not in TITLES:
            raise HTTPException(422, '統計查詢條件不正確')
        for h in hs:
            scoped(u, h)
        first, last = (1,12) if quarter == 0 else ((quarter-1)*3+1, quarter*3)
        start = f'{year}-{first:02}-01'; end = f'{year}-{last:02}-{calendar.monthrange(year,last)[1]}'
        selected_month = __import__('app.excel', fromlist=['period']).period(month) if month is not None else None
        if selected_month:
            year = int(selected_month[:4])
            start = selected_month + '-01'
            end = selected_month + f'-{calendar.monthrange(year,int(selected_month[5:]))[1]:02}'
        today = datetime.now(wf.TZ).date().isoformat()
        with db.connect() as c:
            rows = c.execute('SELECT * FROM monthly_cases WHERE hospital=ANY(%s) AND period>=%s AND period<=%s ORDER BY period,id', (hs,TRACKING_START,f'{year}-12')).fetchall()
            excluded=exclusion_index(c)
            rows=deepcopy(rows)
            for t in rows:
                for x in t['body']['items'].values():x['tracking_excluded_at']=excluded.get((t['hospital'],x['data']['code'],x['data']['name'],x['data']['unit']))
        rows = [t for t in rows if in_domain(t['kind'], domain)]
        if kind:
            rows = [t for t in rows if t['kind'] == kind]
        actual_end = min(end, today)
        output = wf.stats(rows,start,actual_end) if start <= actual_end else None
        quarters = {str(q): wf.stats(rows, f'{year}-{(q-1)*3+1:02}-01', min(today, f'{year}-{q*3:02}-{calendar.monthrange(year,q*3)[1]}')) if f'{year}-{(q-1)*3+1:02}-01' <= today else None for q in range(1,5)}
        return jsonable(dict(year=year, quarter=quarter, month=selected_month, summary=output, quarters=quarters, definition='原因依異常發生月份及期末前已審核的主要原因計算；改善案排除連結提醒。完成率按本期到期案件計算，已完成須經覆核；逾期及進度依期末歷程，缺歷史證據另列。原因季度拆分依本次截止日已審核原因歸屬發生季度；季度變化表回溯當季末進度。金額不是損失或節省。'))

    @app.get('/api/monthly-statistics-export')
    def statistics_export(hospitals:str, year:int, quarter:int=0, kind:str='', u=Depends(auth.current), domain:str='', month:str|None=None):
        data = statistics(hospitals,year,quarter,kind,u,domain,month); s=data['summary']; w=Workbook(); w.remove(w.active)
        sheet(w,'統計口徑',['項目','內容'],[['院別',hospitals],['年度',data['year']],['期間',('本期 '+data['month']) if data.get('month') else ('全年' if not quarter else f'第{quarter}季')],['起始日',s['start'] if s else '尚未開始'],['截止日',s['end'] if s else '尚未開始'],['定義',data['definition']],['歷程不足',s['unknown_history'] if s else '尚未開始']])
        if s:
            sheet(w,'重點指標',['指標','值'],[[label,s[k]] for k,label in [('new_count','新增異常'),('open_count','期末未完成'),('completed_count','已完成件數'),('overdue','逾期'),('due','本期到期'),('completed','到期案已完成'),('ontime','到期案準時完成'),('completion_rate','完成率%'),('ontime_rate','準時完成率%'),('repeated','重複異常品項'),('recurrence','結案後再發生'),('unconfirmed','原因未確認')]])
            sheet(w,'原因分布',['類型','主要原因','件數','需改善','逾期'],[[TITLES[x['kind']],x['cause'],x['count'],x['follow_up'],x['overdue']] for x in s['causes']])
            sheet(w,'季度統計',['季度','新增','期末未完成','逾期','到期','完成率','準時完成率'],[[q,*[v[k] for k in ('new_count','open_count','overdue','due','completion_rate','ontime_rate')]] if v else [q,'尚未開始'] for q,v in data['quarters'].items()])
            sheet(w,'案件明細',['醫院','月份','類型','品號','品名','期末進度','原因','負責人','期限','逾期天數'],[[x[k] for k in ('hospital','period','kind','code','name','status','cause','owner','deadline','overdue')] for x in s['details']])
            sheet(w,'改善前後',['醫院','品號','類型','原始耗用','改善後耗用','原始庫存','改善後庫存','原始週轉天數','改善後週轉天數','改善結果','結果佐證'],[[x['hospital'],x['code'],TITLES[x['kind']],x['before'].get('used'),(x['after'] or {}).get('used'),x['before'].get('ending'),(x['after'] or {}).get('ending'),x['before'].get('days'),(x['after'] or {}).get('days'),x['result'],x['evidence']] for x in s['effects']])
            sheet(w,'審核情形',['指標','值'],[[label,s['review'][k]] for k,label in [('pending','待審核'),('oldest_days','最長等待天數'),('reviewed','本期審核次數'),('returned','本期退回次數'),('return_rate','退回比例%')]])
            sheet(w,'退回原因',['醫院','品號','處理時間','退回原因','審核人'],[[x[k] for k in ('hospital','code','at','reason','reviewer')] for x in s['review']['return_reasons']])
            sheet(w,'負責人待辦',['醫院','負責人','未完成','逾期'],[[x[k] for k in ('hospital','owner','open','overdue')] for x in s['owners']])
            sheet(w,'涉及金額',['項目','金額（非損益）'],[[label,s['amounts'][k]] for k,label in [('usage_increase','耗用增加'),('usage_decrease','耗用減少'),('purchase_increase','單價上漲影響'),('purchase_decrease','單價下降影響'),('inventory','庫存異常涉及金額（跨月觀測累計）')]])
        out=BytesIO(); w.save(out)
        return Response(out.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    @app.get('/api/monthly-cases/{cid}/trend')
    def item_trend(cid:int, item_key:str, month:str, u=Depends(auth.current)):
        p=__import__('app.excel',fromlist=['period']).period(month)
        with db.connect() as c:
            t=c.execute('SELECT * FROM monthly_cases WHERE id=%s',(cid,)).fetchone()
            if not t:raise HTTPException(404,'案件不存在')
            scoped(u,t['hospital']); x=t['body']['items'].get(item_key)
            if not x:raise HTTPException(404,'案件品項不存在')
            periods=[p]
            for _ in range(11):periods.append(engine.previous(periods[-1]))
            points=[]
            for period in reversed(periods):
                d=report(c,t['hospital'],period,baseline='mean3',department='__all__')
                row=next((r for r in d['items'] if key(r)==item_key and r['name']==x['data']['name']),None)
                points.append(dict(period=period,data=snapshot(row,d,t['kind']) if row else None))
        return jsonable(dict(points=points,definition='僅對應同品號品名及單位；缺月不補零；標準化須同範圍實際住院人日。'))

    @app.get('/api/monthly-cause-statistics')
    def cause_statistics(hospitals:str, year:int, u=Depends(auth.current)):
        d = statistics(hospitals,year,0,'',u)
        quarters = {str(q):{cause:sum(v['count'] for v in (d['quarters'][str(q)] or {}).get('causes',[]) if v['cause']==cause) for cause in CAUSES} for q in range(1,5)}
        return dict(year=year, quarters=quarters, annual={cause:sum(q[cause] for q in quarters.values()) for cause in CAUSES}, definition=d['definition'])

    @app.get('/api/monthly-cases/{cid}/history')
    def history(cid:int, u=Depends(auth.current)):
        with db.connect() as c:
            t = c.execute('SELECT hospital FROM monthly_cases WHERE id=%s', (cid,)).fetchone()
            if not t:
                raise HTTPException(404,'案件不存在')
            scoped(u,t['hospital'])
            return jsonable(c.execute("SELECT e.*,u.name operator_name FROM tracking_events e LEFT JOIN users u ON u.id=e.operator WHERE e.body->>'monthly_case_id'=%s ORDER BY e.id DESC", (str(cid),)).fetchall())

    @app.get('/api/monthly-cases-export')
    def export(hospitals:str, month:str, records:bool=False, baseline:str='mean3', view:str='all', keyword:str='', state:str='', cause:str='', date_from:str='', date_to:str='', u=Depends(auth.current), domain:str='', anomaly:str='', mine:bool=False,value_mode:str='amount'):
        if state and state not in wf.STATES:
            raise HTTPException(422, '進度篩選條件不正確')
        result = listing(hospitals,month,records,baseline,u); w = Workbook(); w.remove(w.active); values=[]; events=[]
        in_domain('usage', domain)
        if anomaly and anomaly not in (*TITLES, 'usage_up', 'usage_down'):
            raise HTTPException(422, '異常類型不正確')
        for t in result['cases']:
            if not in_domain(t['kind'], domain):continue
            for k,x in t['items'].items():
                s=x['display_status']
                if not in_view(t, x, view, result['period']):continue
                if mine and not own_todo(x, u):continue
                reading = (x.get('latest_data') or x['data']) if view == 'confirm' else x['data']
                if anomaly in ('usage_up', 'usage_down'):
                    diff = reading.get('difference')
                    if t['kind'] != 'usage' or diff is None or (D(str(diff)) <= 0 if anomaly == 'usage_up' else D(str(diff)) >= 0):continue
                elif anomaly and t['kind'] != anomaly:continue
                if keyword and keyword.lower() not in ' '.join(str(v or '') for v in (x['data']['code'],x['data']['name'])).lower():continue
                if not matches_state(x,state,view):continue
                if cause and x.get('cause_category')!=cause:continue
                for e in x.get('history',[]):
                    if e['action'] in ('approve','approve_follow','finish','reject','reopen','progress','submit','link','unlink') and (not date_from or e['at'][:10]>=date_from) and (not date_to or e['at'][:10]<=date_to):
                        events.append([export_date(e['at']),t['hospital'],t['period'],x['data']['code'],x['data']['name'],e['action'],wf.status(e['snapshot']),e.get('actor'),e['snapshot'].get('review_note') or e['snapshot'].get('progress_note') or e['snapshot'].get('result_note') or e['snapshot'].get('note')])
                d=x['data']; latest=x.get('latest_data') or {}
                values.append([t['hospital'],t['period'],TITLES[t['kind']],d['code'],d['name'],BASES.get(d.get('basis'),'舊版上月'),','.join(d.get('baseline_periods',[])),d['used'],d['difference'],d['pct'],d['ending'],d['days'],d['available'],d.get('standard30'),d.get('standard30_baseline'),s,x.get('cause_category'),x.get('note'),x.get('evidence'),x.get('owner'),x.get('deadline'),x.get('plan'),x.get('target'),x.get('progress_note'),x.get('result_note'),x.get('review_note'),x.get('reviewer'),export_date(x.get('review_date')),x['next_owner'],x['overdue_days'],latest.get('used'),latest.get('ending'),d.get('purchase_pct'),d.get('purchase_impact'),str(x.get('linked_to') or ''),str(x.get('group_root') or ''),reading.get('quantity'),reading.get('baseline_quantity'),reading.get('ending_quantity'),reading.get('pct'),reading.get('baseline'),reading.get('difference')])
        sheet(w,'查詢條件',['醫院','分析年月','檢視比較基準','分頁','分析項目','異常類型','原因','進度','只看我的待辦'],[[hospitals,result['period'],result['baseline_label'],view,domain,anomaly,cause,state,mine]])
        if view=='records':
            sheet(w,'操作紀錄',['處理日期','醫院','原始月份','品號','品名','操作','當時進度','處理人','說明'],sorted(events,key=lambda e:e[0] or datetime.min,reverse=True))
        headers=['醫院','原始月份','異常類型','品號','品名','原始比較基準','原始比較月份','原始耗用金額','原始增減金額','原始增減%','原始庫存金額','原始週轉天數','原始可用天數','原始標準30日每床耗用','基準標準30日每床耗用','處理進度','原因類別','原因說明','佐證','負責人','預計完成日','改善措施','預期結果','最新辦理情形','改善結果','審核說明','審核人','審核日期','下一步負責人','逾期天數','查詢月耗用','查詢月庫存','採購單價增減%','單價變動影響金額','連結原案','同一改善案','本期耗用量','基準耗用量','本期庫存數量','耗用金額變動率（本次比較基準）','基準耗用金額（本次比較基準）','耗用增減金額（本次比較基準）']
        sheet(w,'案件明細',headers,values)
        for ri,row in enumerate(w['案件明細'].iter_rows(min_row=2)):
            for i in (7,8,9,10,11,12,13,14,30,31,32,33,36,37,38,39,40,41):
                val=values[ri][i];row[i].value=float(val) if val is not None else '—';row[i].number_format='0.0"%"' if i in (9,32,39) else '#,##0.0' if i in (11,12,36,37,38) else '#,##0'
            row[27].value=values[ri][27] or '—';row[27].number_format='yyyy-mm-dd hh:mm'
        out=BytesIO();w.save(out)
        return Response(out.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
